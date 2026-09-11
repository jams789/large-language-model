# ============================================================
# Fast Direct VI-RADS ROC Analysis
# Lesion-level + Patient-level
# 7 models, all pairwise comparisons
#
# ============================================================
#
# 分析框架
#
# A. Lesion-level
#
# 每一行 = 1 lesion
#
# Group:
#     lesion-level Group
#
# VI-RADS:
#     lesion-level score
#
# Bootstrap:
#     PatientID 为 cluster
#     同一患者所有 lesions 一起被抽中
#
# 分析：
#     1. Ordinal VI-RADS 1–5 AUC
#     2. VI-RADS >=4 threshold-specific AUC
#     3. VI-RADS >=3 threshold-specific AUC
#
#
# B. Patient-level
#
# 每个 PatientID 聚合成 1 行
#
# Patient Group:
#     max(Group)
#
# Patient VI-RADS:
#     每个评分来源 max(VI-RADS)
#
# Bootstrap:
#     ordinary patient bootstrap
#
# 分析：
#     1. Ordinal VI-RADS 1–5 AUC
#     2. VI-RADS >=4 threshold-specific AUC
#     3. VI-RADS >=3 threshold-specific AUC
#
#
# C. Pairwise comparisons
#
# 7 models -> 21 pairwise comparisons
#
# 对每一种分析：
#
#     Better model = observed AUC 较大的模型
#     Worse model  = observed AUC 较小的模型
#
#     Delta AUC =
#         AUC(Better) - AUC(Worse)
#
# Bootstrap difference:
#
#     Delta_boot =
#         AUC_boot(Better) - AUC_boot(Worse)
#
# 注意：
#     bootstrap 内不重新判断谁大谁小。
#
# 输出：
#
#     Delta AUC
#     bootstrap 95% CI
#     raw two-sided bootstrap P
#     Holm-adjusted P
#
#
# D. Speed optimization
#
# 1. 每个 cohort / level 只生成一次 bootstrap
# 2. 7 个模型共享完全相同 bootstrap samples
# 3. 不再 pair-by-pair bootstrap
# 4. lesion-level 不再反复 pd.concat
# 5. 使用 patient bootstrap counts -> sample weights
# 6. VI-RADS 只有少数离散值，使用快速 weighted AUC
#
#
# 不使用：
#     Logistic regression
#     Youden cutoff
#     DeLong
#
# ============================================================


import os
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from sklearn.metrics import (
    roc_curve,
    confusion_matrix
)

from tkinter import Tk, filedialog


warnings.filterwarnings("ignore")


# ============================================================
# 0. 用户设置
# ============================================================

OUTCOME_COLUMN = "Group"

PATIENT_ID_COLUMN = "PatientID"

COHORT_COLUMN = "cohort"

LESION_ID_COLUMN = "lesion_id"

SURGICAL_COLUMN = "Surgical"

SURGERY_FILTER = "Cystectomy"

SCORE_COLUMNS = [
    "Reference",
    "Qwen3",
    "Original",
    "GPT4",
    "Radiologists",
    "MedGemma",
    "Deepseek"
]


DISPLAY_NAMES = [
    "Reference",
    "Qwen3",
    "Original reports",
    "GPT-4",
    "Radiologists",
    "MedGemma",
    "DeepSeek"
]


COLORS = [
    "#222222",
    "#D55E00",
    "#999999",
    "#CC79A7",
    "#0072B2",
    "#009E73",
    "#E69F00"
]


# ------------------------------------------------------------
# Bootstrap次数
#
# 2000 与你原来的代码一致。
#
# 因为现在共享 bootstrap + 快速 AUC，
# 会远快于逐 pair bootstrap。
# ------------------------------------------------------------

N_BOOTSTRAP = 2000

RANDOM_SEED = 42

DPI = 600


# ------------------------------------------------------------
# Pairwise比较要求完整scores
#
# 推荐保持 True。
#
# 这样7个模型一定基于完全相同病例进行比较，
# paired comparison 最干净。
#
# 如果某个score确实有缺失，会直接报错提醒。
# ------------------------------------------------------------

REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE = True


# ============================================================
# 1. Configuration check
# ============================================================

if len(SCORE_COLUMNS) != len(DISPLAY_NAMES):

    raise ValueError(
        "SCORE_COLUMNS 和 DISPLAY_NAMES 数量必须一致。"
    )


if len(SCORE_COLUMNS) != len(COLORS):

    raise ValueError(
        "SCORE_COLUMNS 和 COLORS 数量必须一致。"
    )


N_MODELS = len(SCORE_COLUMNS)


# ============================================================
# 2. Load data
# ============================================================

def load_data(file_path):

    file_lower = file_path.lower()

    if file_lower.endswith(".xlsx"):

        df = pd.read_excel(
            file_path,
            engine="openpyxl"
        )

    elif file_lower.endswith(".xls"):

        df = pd.read_excel(
            file_path
        )

    elif file_lower.endswith(".csv"):

        try:

            df = pd.read_csv(
                file_path,
                encoding="utf-8-sig"
            )

        except UnicodeDecodeError:

            df = pd.read_csv(
                file_path,
                encoding="gbk"
            )

    else:

        raise ValueError(
            "仅支持 .xlsx / .xls / .csv 文件。"
        )

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    return df


# ============================================================
# 3. Safe divide
# ============================================================

def safe_divide(a, b):

    if b == 0:

        return np.nan

    return a / b


# ============================================================
# 4. Fast weighted AUC
#
# 与 roc_auc_score 对离散 ordinal / binary predictor
# 的定义一致：
#
# AUC =
# P(score_positive > score_negative)
# +
# 0.5 * P(score_positive == score_negative)
#
# 支持 sample weights。
#
# 对 cluster bootstrap：
# 一个患者被抽中 k 次，
# 该患者所有 lesion 的 weight = k。
#
# 避免真正复制 DataFrame。
# ============================================================

def fast_weighted_auc(
    y_true,
    y_score,
    weights=None
):

    y_true = np.asarray(
        y_true,
        dtype=float
    )

    y_score = np.asarray(
        y_score,
        dtype=float
    )

    if weights is None:

        weights = np.ones(
            len(y_true),
            dtype=float
        )

    else:

        weights = np.asarray(
            weights,
            dtype=float
        )

    valid = (
        np.isfinite(y_true)
        &
        np.isfinite(y_score)
        &
        np.isfinite(weights)
        &
        (weights > 0)
    )

    if not np.any(valid):

        return np.nan

    y = y_true[valid].astype(int)

    score = y_score[valid]

    w = weights[valid]

    positive_weight = w[
        y == 1
    ].sum()

    negative_weight = w[
        y == 0
    ].sum()

    if (
        positive_weight <= 0
        or negative_weight <= 0
    ):

        return np.nan

    score_levels = np.sort(
        np.unique(score)
    )

    cumulative_negative = 0.0

    favorable_pairs = 0.0

    for level in score_levels:

        mask_level = (
            score == level
        )

        pos_at_level = w[
            mask_level & (y == 1)
        ].sum()

        neg_at_level = w[
            mask_level & (y == 0)
        ].sum()

        favorable_pairs += (
            pos_at_level
            *
            (
                cumulative_negative
                +
                0.5 * neg_at_level
            )
        )

        cumulative_negative += (
            neg_at_level
        )

    auc = (
        favorable_pairs
        /
        (
            positive_weight
            *
            negative_weight
        )
    )

    return float(auc)


# ============================================================
# 5. Diagnostic metrics
# ============================================================

def diagnostic_metrics(
    y_true,
    score,
    threshold
):

    y_true = np.asarray(
        y_true,
        dtype=int
    )

    score = np.asarray(
        score,
        dtype=float
    )

    valid = (
        np.isfinite(y_true)
        &
        np.isfinite(score)
    )

    y_true = y_true[
        valid
    ]

    score = score[
        valid
    ]

    y_pred = (
        score >= threshold
    ).astype(int)

    cm = confusion_matrix(
        y_true,
        y_pred,
        labels=[
            0,
            1
        ]
    )

    tn, fp, fn, tp = cm.ravel()

    sensitivity = safe_divide(
        tp,
        tp + fn
    )

    specificity = safe_divide(
        tn,
        tn + fp
    )

    ppv = safe_divide(
        tp,
        tp + fp
    )

    npv = safe_divide(
        tn,
        tn + fn
    )

    accuracy = safe_divide(
        tp + tn,
        tp + tn + fp + fn
    )

    balanced_accuracy = (
        sensitivity + specificity
    ) / 2

    return {
        "TP": int(tp),
        "TN": int(tn),
        "FP": int(fp),
        "FN": int(fn),
        "Sensitivity": sensitivity,
        "Specificity": specificity,
        "PPV": ppv,
        "NPV": npv,
        "Accuracy": accuracy,
        "Balanced Accuracy": balanced_accuracy
    }


# ============================================================
# 6. Data check
# ============================================================

def check_dataframe(
    df,
    dataset_name
):

    print("\n" + "=" * 72)

    print(
        f"检查 cohort：{dataset_name}"
    )

    print("=" * 72)

    print(
        f"Lesion 数量：{len(df)}"
    )

    required_columns = [
        OUTCOME_COLUMN,
        PATIENT_ID_COLUMN,
        LESION_ID_COLUMN
    ] + SCORE_COLUMNS

    missing_columns = [
        col
        for col in required_columns
        if col not in df.columns
    ]

    if missing_columns:

        raise ValueError(
            f"\n{dataset_name} 缺少以下列：\n"
            f"{missing_columns}"
        )

    n_patients = (
        df[PATIENT_ID_COLUMN]
        .dropna()
        .nunique()
    )

    print(
        f"Patient 数量：{n_patients}"
    )

    y = pd.to_numeric(
        df[OUTCOME_COLUMN],
        errors="coerce"
    )

    unique_y = sorted(
        y.dropna()
        .unique()
        .tolist()
    )

    print(
        f"{OUTCOME_COLUMN} 唯一值：{unique_y}"
    )

    if not set(unique_y).issubset(
        {0, 1}
    ):

        raise ValueError(
            f"{OUTCOME_COLUMN} 必须为 0/1。"
        )

    if len(unique_y) < 2:

        raise ValueError(
            f"{dataset_name} 中没有同时包含 0 和 1。"
        )

    print(
        "\nLesion-level Group distribution:"
    )

    print(
        y.value_counts(
            dropna=False
        ).sort_index()
    )

    print(
        "\nVI-RADS score ranges:"
    )

    score_missing = {}

    for score_col in SCORE_COLUMNS:

        values = pd.to_numeric(
            df[score_col],
            errors="coerce"
        )

        valid = values.dropna()

        score_missing[
            score_col
        ] = int(
            values.isna().sum()
        )

        if len(valid) == 0:

            print(
                f"{score_col:<20}: 全部缺失"
            )

            continue

        print(
            f"{score_col:<20}: "
            f"min={valid.min():.0f}, "
            f"max={valid.max():.0f}, "
            f"N={len(valid)}, "
            f"missing={values.isna().sum()}"
        )

        invalid = valid[
            ~valid.isin(
                [1, 2, 3, 4, 5]
            )
        ]

        if len(invalid) > 0:

            raise ValueError(
                f"{score_col} 中存在非1–5评分："
                f"{sorted(invalid.unique())}"
            )

    if REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE:

        incomplete = {
            k: v
            for k, v in score_missing.items()
            if v > 0
        }

        if len(incomplete) > 0:

            raise ValueError(
                "\nPairwise paired comparison 要求7个score"
                "在当前cohort中均完整。\n"
                f"发现缺失：{incomplete}\n"
                "请先检查数据，或将 "
                "REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE=False。"
            )


# ============================================================
# 7. Build patient-level data
# ============================================================

def build_patient_level_data(
    df,
    dataset_name=""
):

    required_columns = [
        PATIENT_ID_COLUMN,
        OUTCOME_COLUMN
    ] + SCORE_COLUMNS

    temp = df[
        required_columns
    ].copy()

    temp[OUTCOME_COLUMN] = pd.to_numeric(
        temp[OUTCOME_COLUMN],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        temp[score_col] = pd.to_numeric(
            temp[score_col],
            errors="coerce"
        )

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    )

    if len(temp) == 0:

        return pd.DataFrame(
            columns=required_columns
        )

    group_nunique = (
        temp
        .groupby(
            PATIENT_ID_COLUMN
        )[OUTCOME_COLUMN]
        .nunique()
    )

    discordant_patients = (
        group_nunique[
            group_nunique > 1
        ]
        .index
        .tolist()
    )

    print("\n" + "-" * 72)

    print(
        f"{dataset_name}: 构建 patient-level 数据"
    )

    print("-" * 72)

    if len(discordant_patients) > 0:

        print(
            f"发现 {len(discordant_patients)} 名患者"
            "同时具有 Group=0 和 Group=1 lesions。"
        )

        print(
            "Patient-level Group = max(Group)。"
        )

    aggregation_dict = {
        OUTCOME_COLUMN: "max"
    }

    for score_col in SCORE_COLUMNS:

        aggregation_dict[
            score_col
        ] = "max"

    patient_df = (
        temp
        .groupby(
            PATIENT_ID_COLUMN,
            as_index=False
        )
        .agg(
            aggregation_dict
        )
    )

    patient_df[
        OUTCOME_COLUMN
    ] = patient_df[
        OUTCOME_COLUMN
    ].astype(int)

    print(
        f"Patient 数量：{len(patient_df)}"
    )

    print(
        "Patient-level pathology distribution:"
    )

    print(
        patient_df[
            OUTCOME_COLUMN
        ]
        .value_counts()
        .sort_index()
    )

    return patient_df


# ============================================================
# 8. Bootstrap matrix
#
# 返回：
#
# point_auc:
#     shape = (3, 7)
#
# bootstrap_auc:
#     shape = (N_BOOTSTRAP, 3, 7)
#
# analysis index:
#
#     0 = Ordinal
#     1 = >=4
#     2 = >=3
#
# ============================================================

ANALYSIS_NAMES = [
    "Ordinal",
    ">=4",
    ">=3"
]


def transform_score(
    score,
    analysis_index
):

    if analysis_index == 0:

        return score

    elif analysis_index == 1:

        return (
            score >= 4
        ).astype(float)

    elif analysis_index == 2:

        return (
            score >= 3
        ).astype(float)

    else:

        raise ValueError(
            "Unknown analysis index."
        )


# ============================================================
# 9. Lesion-level bootstrap matrix
#
# Bootstrap cluster = PatientID
#
# Multinomial counts 与：
#
# np.random.choice(
#     patients,
#     size=n_patients,
#     replace=True
# )
#
# 完全等价。
#
# 若 patient 被抽中 k 次：
# 其所有 lesion weight = k
# ============================================================

def bootstrap_matrix_lesion(
    df,
    n_bootstrap=N_BOOTSTRAP,
    seed=RANDOM_SEED
):

    temp = df[
        [
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ] + SCORE_COLUMNS
    ].copy()

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).reset_index(
        drop=True
    )

    y = temp[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    scores = (
        temp[
            SCORE_COLUMNS
        ]
        .astype(float)
        .to_numpy()
    )

    patient_codes, unique_patients = pd.factorize(
        temp[
            PATIENT_ID_COLUMN
        ],
        sort=False
    )

    n_patients = len(
        unique_patients
    )

    n_rows = len(temp)

    print(
        f"\nLesion bootstrap: "
        f"{n_rows} lesions / "
        f"{n_patients} patients / "
        f"{n_bootstrap} replicates"
    )

    point_auc = np.full(
        (
            3,
            N_MODELS
        ),
        np.nan
    )

    bootstrap_auc = np.full(
        (
            n_bootstrap,
            3,
            N_MODELS
        ),
        np.nan,
        dtype=float
    )

    # --------------------------------------------------------
    # observed AUC
    # --------------------------------------------------------

    for analysis_idx in range(3):

        for model_idx in range(
            N_MODELS
        ):

            predictor = transform_score(
                scores[
                    :,
                    model_idx
                ],
                analysis_idx
            )

            point_auc[
                analysis_idx,
                model_idx
            ] = fast_weighted_auc(
                y,
                predictor
            )

    # --------------------------------------------------------
    # bootstrap
    # --------------------------------------------------------

    rng = np.random.default_rng(
        seed
    )

    probabilities = np.full(
        n_patients,
        1.0 / n_patients
    )

    progress_marks = set(
        [
            int(
                n_bootstrap * x
            )
            for x in [
                0.25,
                0.50,
                0.75,
                1.00
            ]
        ]
    )

    for b in range(
        n_bootstrap
    ):

        patient_counts = rng.multinomial(
            n_patients,
            probabilities
        )

        row_weights = patient_counts[
            patient_codes
        ].astype(float)

        for model_idx in range(
            N_MODELS
        ):

            score_original = scores[
                :,
                model_idx
            ]

            # ordinal

            bootstrap_auc[
                b,
                0,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_original,
                row_weights
            )

            # >=4

            score_ge4 = (
                score_original >= 4
            ).astype(float)

            bootstrap_auc[
                b,
                1,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge4,
                row_weights
            )

            # >=3

            score_ge3 = (
                score_original >= 3
            ).astype(float)

            bootstrap_auc[
                b,
                2,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge3,
                row_weights
            )

        current = b + 1

        if current in progress_marks:

            pct = (
                100 * current / n_bootstrap
            )

            print(
                f"  Lesion bootstrap: "
                f"{pct:.0f}%"
            )

    return (
        point_auc,
        bootstrap_auc
    )


# ============================================================
# 10. Patient-level bootstrap matrix
# ============================================================

def bootstrap_matrix_patient(
    patient_df,
    n_bootstrap=N_BOOTSTRAP,
    seed=RANDOM_SEED
):

    temp = patient_df[
        [
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ] + SCORE_COLUMNS
    ].copy()

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).reset_index(
        drop=True
    )

    y = temp[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    scores = (
        temp[
            SCORE_COLUMNS
        ]
        .astype(float)
        .to_numpy()
    )

    n_patients = len(temp)

    print(
        f"\nPatient bootstrap: "
        f"{n_patients} patients / "
        f"{n_bootstrap} replicates"
    )

    point_auc = np.full(
        (
            3,
            N_MODELS
        ),
        np.nan
    )

    bootstrap_auc = np.full(
        (
            n_bootstrap,
            3,
            N_MODELS
        ),
        np.nan,
        dtype=float
    )

    # --------------------------------------------------------
    # observed
    # --------------------------------------------------------

    for analysis_idx in range(3):

        for model_idx in range(
            N_MODELS
        ):

            predictor = transform_score(
                scores[
                    :,
                    model_idx
                ],
                analysis_idx
            )

            point_auc[
                analysis_idx,
                model_idx
            ] = fast_weighted_auc(
                y,
                predictor
            )

    # --------------------------------------------------------
    # bootstrap
    # --------------------------------------------------------

    rng = np.random.default_rng(
        seed
    )

    probabilities = np.full(
        n_patients,
        1.0 / n_patients
    )

    progress_marks = set(
        [
            int(
                n_bootstrap * x
            )
            for x in [
                0.25,
                0.50,
                0.75,
                1.00
            ]
        ]
    )

    for b in range(
        n_bootstrap
    ):

        patient_weights = rng.multinomial(
            n_patients,
            probabilities
        ).astype(float)

        for model_idx in range(
            N_MODELS
        ):

            score_original = scores[
                :,
                model_idx
            ]

            bootstrap_auc[
                b,
                0,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_original,
                patient_weights
            )

            score_ge4 = (
                score_original >= 4
            ).astype(float)

            bootstrap_auc[
                b,
                1,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge4,
                patient_weights
            )

            score_ge3 = (
                score_original >= 3
            ).astype(float)

            bootstrap_auc[
                b,
                2,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge3,
                patient_weights
            )

        current = b + 1

        if current in progress_marks:

            pct = (
                100 * current / n_bootstrap
            )

            print(
                f"  Patient bootstrap: "
                f"{pct:.0f}%"
            )

    return (
        point_auc,
        bootstrap_auc
    )


# ============================================================
# 11. Bootstrap CI
# ============================================================

def bootstrap_ci(
    bootstrap_values
):

    values = np.asarray(
        bootstrap_values,
        dtype=float
    )

    values = values[
        np.isfinite(values)
    ]

    if len(values) == 0:

        return (
            np.nan,
            np.nan
        )

    lower = np.percentile(
        values,
        2.5
    )

    upper = np.percentile(
        values,
        97.5
    )

    return (
        float(lower),
        float(upper)
    )


# ============================================================
# 12. Two-sided paired bootstrap P
#
# observed better-worse orientation固定。
#
# P =
# 2 * min[
#     Pr(Delta_boot <= 0),
#     Pr(Delta_boot >= 0)
# ]
#
# +1 correction 防止出现P=0。
# ============================================================

def paired_bootstrap_pvalue(
    delta_boot
):

    delta_boot = np.asarray(
        delta_boot,
        dtype=float
    )

    delta_boot = delta_boot[
        np.isfinite(delta_boot)
    ]

    n = len(
        delta_boot
    )

    if n == 0:

        return np.nan

    n_le_zero = np.sum(
        delta_boot <= 0
    )

    n_ge_zero = np.sum(
        delta_boot >= 0
    )

    p_lower = (
        n_le_zero + 1
    ) / (
        n + 1
    )

    p_upper = (
        n_ge_zero + 1
    ) / (
        n + 1
    )

    p = 2 * min(
        p_lower,
        p_upper
    )

    return min(
        float(p),
        1.0
    )


# ============================================================
# 13. Holm correction
# ============================================================

def holm_adjust(
    p_values
):

    p_values = np.asarray(
        p_values,
        dtype=float
    )

    adjusted = np.full(
        len(p_values),
        np.nan
    )

    valid_idx = np.where(
        np.isfinite(
            p_values
        )
    )[0]

    if len(valid_idx) == 0:

        return adjusted

    valid_p = p_values[
        valid_idx
    ]

    order = np.argsort(
        valid_p
    )

    sorted_p = valid_p[
        order
    ]

    m = len(
        sorted_p
    )

    adjusted_sorted = np.zeros(
        m,
        dtype=float
    )

    running_max = 0.0

    for i in range(
        m
    ):

        value = (
            (m - i)
            *
            sorted_p[i]
        )

        value = min(
            value,
            1.0
        )

        running_max = max(
            running_max,
            value
        )

        adjusted_sorted[i] = (
            running_max
        )

    reverse_order = np.empty(
        m,
        dtype=int
    )

    reverse_order[
        order
    ] = np.arange(
        m
    )

    adjusted_valid = adjusted_sorted[
        reverse_order
    ]

    adjusted[
        valid_idx
    ] = adjusted_valid

    return adjusted


# ============================================================
# 14. Build individual AUC table
# ============================================================

def build_auc_table(
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc,
    n_lesions=None,
    n_patients=None
):

    rows = []

    if level_name == "Lesion":

        ci_method = (
            "Patient-level cluster bootstrap"
        )

    else:

        ci_method = (
            "Patient-level bootstrap"
        )

    for analysis_idx, analysis_name in enumerate(
        ANALYSIS_NAMES
    ):

        for model_idx, score_col in enumerate(
            SCORE_COLUMNS
        ):

            auc_value = point_auc[
                analysis_idx,
                model_idx
            ]

            lower, upper = bootstrap_ci(
                bootstrap_auc[
                    :,
                    analysis_idx,
                    model_idx
                ]
            )

            row = {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    level_name,

                "Analysis":
                    analysis_name,

                "Variable":
                    score_col,

                "Model":
                    DISPLAY_NAMES[
                        model_idx
                    ],

                "AUC":
                    auc_value,

                "95% CI lower":
                    lower,

                "95% CI upper":
                    upper,

                "AUC (95% CI)":
                    (
                        f"{auc_value:.3f} "
                        f"({lower:.3f}–{upper:.3f})"
                    ),

                "CI method":
                    ci_method,

                "Bootstrap N":
                    N_BOOTSTRAP
            }

            if n_lesions is not None:

                row[
                    "Lesions"
                ] = n_lesions

            if n_patients is not None:

                row[
                    "Patients"
                ] = n_patients

            rows.append(
                row
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# 15. Build pairwise comparison table
# ============================================================

def build_pairwise_table(
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc
):

    all_rows = []

    for analysis_idx, analysis_name in enumerate(
        ANALYSIS_NAMES
    ):

        current_rows = []

        for i in range(
            N_MODELS
        ):

            for j in range(
                i + 1,
                N_MODELS
            ):

                auc_i = point_auc[
                    analysis_idx,
                    i
                ]

                auc_j = point_auc[
                    analysis_idx,
                    j
                ]

                if (
                    not np.isfinite(
                        auc_i
                    )
                    or
                    not np.isfinite(
                        auc_j
                    )
                ):

                    continue

                # --------------------------------------------
                # observed AUC determines direction
                # --------------------------------------------

                if auc_i >= auc_j:

                    better_idx = i

                    worse_idx = j

                else:

                    better_idx = j

                    worse_idx = i

                auc_better = point_auc[
                    analysis_idx,
                    better_idx
                ]

                auc_worse = point_auc[
                    analysis_idx,
                    worse_idx
                ]

                delta_auc = (
                    auc_better
                    -
                    auc_worse
                )

                delta_boot = (
                    bootstrap_auc[
                        :,
                        analysis_idx,
                        better_idx
                    ]
                    -
                    bootstrap_auc[
                        :,
                        analysis_idx,
                        worse_idx
                    ]
                )

                valid_delta = delta_boot[
                    np.isfinite(
                        delta_boot
                    )
                ]

                delta_lower, delta_upper = (
                    bootstrap_ci(
                        valid_delta
                    )
                )

                p_raw = paired_bootstrap_pvalue(
                    valid_delta
                )

                current_rows.append(
                    {
                        "Dataset":
                            dataset_name,

                        "Analysis level":
                            level_name,

                        "Analysis":
                            analysis_name,

                        "Model 1":
                            DISPLAY_NAMES[i],

                        "Model 2":
                            DISPLAY_NAMES[j],

                        "Better model":
                            DISPLAY_NAMES[
                                better_idx
                            ],

                        "Worse model":
                            DISPLAY_NAMES[
                                worse_idx
                            ],

                        "AUC better":
                            auc_better,

                        "AUC worse":
                            auc_worse,

                        "Delta AUC":
                            delta_auc,

                        "Delta 95% CI lower":
                            delta_lower,

                        "Delta 95% CI upper":
                            delta_upper,

                        "Delta AUC (95% CI)":
                            (
                                f"{delta_auc:.3f} "
                                f"({delta_lower:.3f}–"
                                f"{delta_upper:.3f})"
                            ),

                        "P raw":
                            p_raw,

                        "Valid bootstrap replicates":
                            len(
                                valid_delta
                            ),

                        "Comparison method":
                            (
                                "Paired bootstrap; "
                                "shared bootstrap samples"
                            )
                    }
                )

        # ----------------------------------------------------
        # Holm correction:
        # 每个 cohort × level × analysis
        # 的21组comparison作为一个family
        # ----------------------------------------------------

        if len(
            current_rows
        ) > 0:

            raw_ps = [
                row[
                    "P raw"
                ]
                for row in current_rows
            ]

            adjusted_ps = holm_adjust(
                raw_ps
            )

            for row, p_adj in zip(
                current_rows,
                adjusted_ps
            ):

                row[
                    "P Holm"
                ] = p_adj

                row[
                    "Significant raw P<0.05"
                ] = (
                    pd.notna(
                        row[
                            "P raw"
                        ]
                    )
                    and
                    row[
                        "P raw"
                    ] < 0.05
                )

                row[
                    "Significant Holm P<0.05"
                ] = (
                    pd.notna(
                        p_adj
                    )
                    and
                    p_adj < 0.05
                )

            all_rows.extend(
                current_rows
            )

    return pd.DataFrame(
        all_rows
    )


# ============================================================
# 16. Threshold diagnostic table
# ============================================================

def build_threshold_metrics_table(
    source_df,
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc
):

    rows = []

    if level_name == "Lesion":

        n_lesions = len(
            source_df
        )

        n_patients = (
            source_df[
                PATIENT_ID_COLUMN
            ]
            .nunique()
        )

        ci_method = (
            "Patient-level cluster bootstrap"
        )

    else:

        n_lesions = None

        n_patients = len(
            source_df
        )

        ci_method = (
            "Patient-level bootstrap"
        )

    y = source_df[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    for model_idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        score = source_df[
            score_col
        ].astype(float).to_numpy()

        for threshold, analysis_idx in [
            (
                4,
                1
            ),
            (
                3,
                2
            )
        ]:

            metrics = diagnostic_metrics(
                y,
                score,
                threshold
            )

            auc_value = point_auc[
                analysis_idx,
                model_idx
            ]

            lower, upper = bootstrap_ci(
                bootstrap_auc[
                    :,
                    analysis_idx,
                    model_idx
                ]
            )

            row = {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    level_name,

                "Variable":
                    score_col,

                "Model":
                    DISPLAY_NAMES[
                        model_idx
                    ],

                "Threshold":
                    f">={threshold}",

                "Threshold-specific AUC":
                    auc_value,

                "AUC 95% CI lower":
                    lower,

                "AUC 95% CI upper":
                    upper,

                "Threshold AUC (95% CI)":
                    (
                        f"{auc_value:.3f} "
                        f"({lower:.3f}–{upper:.3f})"
                    ),

                "TP":
                    metrics[
                        "TP"
                    ],

                "TN":
                    metrics[
                        "TN"
                    ],

                "FP":
                    metrics[
                        "FP"
                    ],

                "FN":
                    metrics[
                        "FN"
                    ],

                "Sensitivity":
                    metrics[
                        "Sensitivity"
                    ],

                "Specificity":
                    metrics[
                        "Specificity"
                    ],

                "PPV":
                    metrics[
                        "PPV"
                    ],

                "NPV":
                    metrics[
                        "NPV"
                    ],

                "Accuracy":
                    metrics[
                        "Accuracy"
                    ],

                "Balanced Accuracy":
                    metrics[
                        "Balanced Accuracy"
                    ],

                "AUC CI method":
                    ci_method
            }

            if n_lesions is not None:

                row[
                    "Lesions"
                ] = n_lesions

            row[
                "Patients"
            ] = n_patients

            rows.append(
                row
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# 17. Plot style
# ============================================================

def set_plot_style():

    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 11,
            "axes.labelsize": 14,
            "axes.titlesize": 15,
            "axes.linewidth": 1.2,
            "xtick.labelsize": 11,
            "ytick.labelsize": 11,
            "xtick.major.size": 5,
            "ytick.major.size": 5,
            "xtick.major.width": 1.1,
            "ytick.major.width": 1.1,
            "legend.fontsize": 8.5,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none"
        }
    )


set_plot_style()


# ============================================================
# 18. ROC plot
#
# 只画 ordinal VI-RADS 1–5 ROC
# 与你原来的图保持一致
# ============================================================

def plot_ordinal_roc(
    source_df,
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc,
    output_dir
):

    plt.figure(
        figsize=(
            8,
            6
        )
    )

    y = source_df[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    for model_idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        score = source_df[
            score_col
        ].astype(float).to_numpy()

        valid = (
            np.isfinite(y)
            &
            np.isfinite(score)
        )

        y_current = y[
            valid
        ]

        score_current = score[
            valid
        ]

        if np.unique(
            y_current
        ).size < 2:

            continue

        fpr, tpr, _ = roc_curve(
            y_current,
            score_current
        )

        auc_value = point_auc[
            0,
            model_idx
        ]

        lower, upper = bootstrap_ci(
            bootstrap_auc[
                :,
                0,
                model_idx
            ]
        )

        label_text = (
            f"{DISPLAY_NAMES[model_idx]} "
            f"(AUC = {auc_value:.3f} "
            f"[{lower:.3f}-{upper:.3f}])"
        )

        plt.plot(
            fpr,
            tpr,
            color=COLORS[
                model_idx
            ],
            linewidth=2,
            label=label_text
        )

    plt.plot(
        [
            0,
            1
        ],
        [
            0,
            1
        ],
        color="navy",
        linestyle="--",
        linewidth=1.2
    )

    plt.xlim(
        [
            -0.04,
            1.04
        ]
    )

    plt.ylim(
        [
            -0.04,
            1.04
        ]
    )

    plt.xlabel(
        "1 - Specificity"
    )

    plt.ylabel(
        "Sensitivity"
    )

    plt.title(
        f"{level_name}-level ROC - "
        f"{dataset_name}"
    )

    plt.legend(
        loc="lower right"
    )

    plt.grid(
        False
    )

    safe_name = (
        str(
            dataset_name
        )
        .replace(
            " ",
            "_"
        )
        .replace(
            "/",
            "_"
        )
        .replace(
            "\\",
            "_"
        )
        .replace(
            ":",
            "_"
        )
    )

    for ext in [
        "pdf",
        "svg",
        "png"
    ]:

        file_name = os.path.join(
            output_dir,
            (
                f"{safe_name}-"
                f"{level_name}-level-ROC."
                f"{ext}"
            )
        )

        if ext == "png":

            plt.savefig(
                file_name,
                dpi=DPI,
                bbox_inches="tight"
            )

        else:

            plt.savefig(
                file_name,
                format=ext,
                bbox_inches="tight"
            )

    plt.close()


# ============================================================
# 19. Analyze one cohort
# ============================================================

def analyze_cohort(
    df_current,
    dataset_name,
    output_dir
):

    print("\n" + "#" * 72)

    print(
        f"COHORT: {dataset_name}"
    )

    print("#" * 72)

    check_dataframe(
        df_current,
        dataset_name
    )

    # ========================================================
    # Lesion level
    # ========================================================

    print("\n" + "=" * 72)

    print(
        f"LESION-LEVEL: {dataset_name}"
    )

    print("=" * 72)

    (
        lesion_point_auc,
        lesion_boot_auc
    ) = bootstrap_matrix_lesion(
        df_current,
        n_bootstrap=N_BOOTSTRAP,
        seed=RANDOM_SEED
    )

    lesion_auc_table = build_auc_table(
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc,
        n_lesions=len(
            df_current
        ),
        n_patients=df_current[
            PATIENT_ID_COLUMN
        ].nunique()
    )

    lesion_pairwise = build_pairwise_table(
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc
    )

    lesion_threshold = (
        build_threshold_metrics_table(
            source_df=df_current,
            dataset_name=dataset_name,
            level_name="Lesion",
            point_auc=lesion_point_auc,
            bootstrap_auc=lesion_boot_auc
        )
    )

    plot_ordinal_roc(
        source_df=df_current,
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc,
        output_dir=output_dir
    )

    # ========================================================
    # Patient level
    # ========================================================

    print("\n" + "=" * 72)

    print(
        f"PATIENT-LEVEL: {dataset_name}"
    )

    print("=" * 72)

    patient_df = build_patient_level_data(
        df_current,
        dataset_name
    )

    (
        patient_point_auc,
        patient_boot_auc
    ) = bootstrap_matrix_patient(
        patient_df,
        n_bootstrap=N_BOOTSTRAP,
        seed=RANDOM_SEED + 100000
    )

    patient_auc_table = build_auc_table(
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc,
        n_patients=len(
            patient_df
        )
    )

    patient_pairwise = build_pairwise_table(
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc
    )

    patient_threshold = (
        build_threshold_metrics_table(
            source_df=patient_df,
            dataset_name=dataset_name,
            level_name="Patient",
            point_auc=patient_point_auc,
            bootstrap_auc=patient_boot_auc
        )
    )

    plot_ordinal_roc(
        source_df=patient_df,
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc,
        output_dir=output_dir
    )

    return {
        "lesion_auc":
            lesion_auc_table,

        "lesion_threshold":
            lesion_threshold,

        "lesion_pairwise":
            lesion_pairwise,

        "patient_auc":
            patient_auc_table,

        "patient_threshold":
            patient_threshold,

        "patient_pairwise":
            patient_pairwise
    }


# ============================================================
# 20. Excel formatting
# ============================================================

def format_excel(
    excel_file
):

    try:

        from openpyxl import (
            load_workbook
        )

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment
        )

        from openpyxl.utils import (
            get_column_letter
        )

        wb = load_workbook(
            excel_file
        )

        decimal_columns = [
            "AUC",
            "95% CI lower",
            "95% CI upper",
            "Threshold-specific AUC",
            "AUC 95% CI lower",
            "AUC 95% CI upper",
            "Sensitivity",
            "Specificity",
            "PPV",
            "NPV",
            "Accuracy",
            "Balanced Accuracy",
            "AUC better",
            "AUC worse",
            "Delta AUC",
            "Delta 95% CI lower",
            "Delta 95% CI upper",
            "P raw",
            "P Holm"
        ]

        for ws in wb.worksheets:

            if ws.max_row < 1:

                continue

            for cell in ws[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor="4F81BD"
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            header_map = {
                cell.value:
                    cell.column
                for cell in ws[1]
            }

            for column_cells in ws.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[
                            0
                        ].column
                    )
                )

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value is None
                        else str(
                            cell.value
                        )
                    )

                    max_length = max(
                        max_length,
                        len(
                            value
                        )
                    )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    42
                )

            ws.freeze_panes = (
                "A2"
            )

            ws.auto_filter.ref = (
                ws.dimensions
            )

            for col_name in decimal_columns:

                if col_name not in header_map:

                    continue

                col_idx = header_map[
                    col_name
                ]

                for row_idx in range(
                    2,
                    ws.max_row + 1
                ):

                    ws.cell(
                        row=row_idx,
                        column=col_idx
                    ).number_format = (
                        "0.000"
                    )

        wb.save(
            excel_file
        )

    except Exception as e:

        print(
            f"Excel 格式化跳过：{e}"
        )


# ============================================================
# 21. Safe concat
# ============================================================

def safe_concat(
    list_of_dfs
):

    valid = [
        x
        for x in list_of_dfs
        if (
            x is not None
            and len(x) > 0
        )
    ]

    if len(valid) == 0:

        return pd.DataFrame()

    return pd.concat(
        valid,
        ignore_index=True
    )


# ============================================================
# 22. Main
# ============================================================

def main():

    root = Tk()

    root.withdraw()

    try:

        root.attributes(
            "-topmost",
            True
        )

    except Exception:

        pass

    print("\n" + "=" * 72)

    print(
        "FAST VI-RADS ROC + ALL PAIRWISE AUC COMPARISONS"
    )

    print("=" * 72)

    print(
        f"\nModels: {N_MODELS}"
    )

    print(
        "Pairwise comparisons per analysis: 21"
    )

    print(
        f"Bootstrap replicates: {N_BOOTSTRAP}"
    )

    print(
        "\nAnalysis:"
    )

    print(
        "1. Lesion ordinal"
    )

    print(
        "2. Lesion >=4"
    )

    print(
        "3. Lesion >=3"
    )

    print(
        "4. Patient ordinal"
    )

    print(
        "5. Patient >=4"
    )

    print(
        "6. Patient >=3"
    )

    file_path = (
        filedialog.askopenfilename(
            title=(
                "选择 VI-RADS 数据文件"
            ),
            filetypes=[
                (
                    "Excel / CSV",
                    "*.xlsx *.xls *.csv"
                ),
                (
                    "All files",
                    "*.*"
                )
            ]
        )
    )

    root.destroy()

    if not file_path:

        print(
            "没有选择文件。"
        )

        return

    df = load_data(
        file_path
    )

    print(
        f"\n读取文件：{file_path}"
    )

    print(
        f"数据维度：{df.shape}"
    )

    # --------------------------------------------------------
    # required columns
    # --------------------------------------------------------

    required = [
    COHORT_COLUMN,
    PATIENT_ID_COLUMN,
    LESION_ID_COLUMN,
    SURGICAL_COLUMN,
    OUTCOME_COLUMN
] + SCORE_COLUMNS

    missing = [
        col
        for col in required
        if col not in df.columns
    ]

    if missing:

        raise ValueError(
            f"缺少列：{missing}"
        )

    # --------------------------------------------------------
    # cohort cleanup
    # --------------------------------------------------------

    mask = df[
        COHORT_COLUMN
    ].notna()

    df.loc[
        mask,
        COHORT_COLUMN
    ] = (
        df.loc[
            mask,
            COHORT_COLUMN
        ]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    df.loc[
        df[
            COHORT_COLUMN
        ] == "",
        COHORT_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # PatientID cleanup
    # --------------------------------------------------------

    mask = df[
        PATIENT_ID_COLUMN
    ].notna()

    df.loc[
        mask,
        PATIENT_ID_COLUMN
    ] = (
        df.loc[
            mask,
            PATIENT_ID_COLUMN
        ]
        .astype(str)
        .str.strip()
    )

    df.loc[
        df[
            PATIENT_ID_COLUMN
        ] == "",
        PATIENT_ID_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # lesion_id cleanup
    # --------------------------------------------------------

    mask = df[
        LESION_ID_COLUMN
    ].notna()

    df.loc[
        mask,
        LESION_ID_COLUMN
    ] = (
        df.loc[
            mask,
            LESION_ID_COLUMN
        ]
        .astype(str)
        .str.strip()
    )
    # --------------------------------------------------------
    # Surgical cleanup
    # --------------------------------------------------------

    mask = df[
        SURGICAL_COLUMN
    ].notna()

    df.loc[
        mask,
        SURGICAL_COLUMN
    ] = (
        df.loc[
            mask,
            SURGICAL_COLUMN
        ]
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # numeric conversion
    # --------------------------------------------------------

    df[
        OUTCOME_COLUMN
    ] = pd.to_numeric(
        df[
            OUTCOME_COLUMN
        ],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        df[
            score_col
        ] = pd.to_numeric(
            df[
                score_col
            ],
            errors="coerce"
        )

    # --------------------------------------------------------
    # Remove invalid patient / outcome
    # --------------------------------------------------------

    before = len(
        df
    )

    df = df.dropna(
        subset=[
            COHORT_COLUMN,
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).copy()

    after = len(
        df
    )

    if before != after:

        print(
            f"\n删除 {before-after} 行 "
            "cohort / PatientID / Group 缺失记录。"
        )

    # --------------------------------------------------------
    # cohorts
    # --------------------------------------------------------

    cohort_names = (
        df[
            COHORT_COLUMN
        ]
        .dropna()
        .unique()
        .tolist()
    )

    if len(cohort_names) == 0:

        raise ValueError(
            "cohort列没有有效值。"
        )

    print(
        "\nCohorts:"
    )

    print(
        df[
            COHORT_COLUMN
        ]
        .value_counts()
    )

    # --------------------------------------------------------
    # output dir
    # --------------------------------------------------------

    input_dir = os.path.dirname(
        os.path.abspath(
            file_path
        )
    )

    output_dir = os.path.join(
        input_dir,
        "Fast_VIRADS_Cystectomy_Sensitivity"
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    all_lesion_auc = []

    all_lesion_threshold = []

    all_lesion_pairwise = []

    all_patient_auc = []

    all_patient_threshold = []

    all_patient_pairwise = []

    # ========================================================
    # cohort loop
    # ========================================================

    for cohort_number, cohort_name in enumerate(
        cohort_names
    ):

        df_current = df[
            (
                df[COHORT_COLUMN] == cohort_name
            )
            &
            (
                df[SURGICAL_COLUMN] == SURGERY_FILTER
            )
        ].copy()

        if len(df_current) == 0:

            print(
                f"\n{cohort_name}: 没有 {SURGERY_FILTER} 病例，跳过。"
            )

            continue

        results = analyze_cohort(
            df_current=df_current,
            dataset_name=f"{cohort_name}_{SURGERY_FILTER}",
            output_dir=output_dir
        )

        all_lesion_auc.append(
            results["lesion_auc"]
        )

        all_lesion_threshold.append(
            results["lesion_threshold"]
        )

        all_lesion_pairwise.append(
            results["lesion_pairwise"]
        )

        all_patient_auc.append(
            results["patient_auc"]
        )

        all_patient_threshold.append(
            results["patient_threshold"]
        )

        all_patient_pairwise.append(
            results["patient_pairwise"]
        )

    # ========================================================
    # combine
    # ========================================================

    lesion_auc = safe_concat(
        all_lesion_auc
    )

    lesion_threshold = safe_concat(
        all_lesion_threshold
    )

    lesion_pairwise = safe_concat(
        all_lesion_pairwise
    )

    patient_auc = safe_concat(
        all_patient_auc
    )

    patient_threshold = safe_concat(
        all_patient_threshold
    )

    patient_pairwise = safe_concat(
        all_patient_pairwise
    )

    # ========================================================
    # Split tables
    # ========================================================

    lesion_ordinal_auc = (
        lesion_auc[
            lesion_auc[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ordinal_auc = (
        patient_auc[
            patient_auc[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_ge4 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_ge3 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge4 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge3 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ordinal = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ge4 = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ge3 = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ordinal = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ge4 = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ge3 = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    # ========================================================
    # Excel
    # ========================================================

    excel_file = os.path.join(
        output_dir,
        "VI-RADS_FAST_all_pairwise_analysis.xlsx"
    )

    with pd.ExcelWriter(
        excel_file,
        engine="openpyxl"
    ) as writer:

        # individual performance

        lesion_ordinal_auc.to_excel(
            writer,
            sheet_name="Lesion_AUC",
            index=False
        )

        lesion_ge4.to_excel(
            writer,
            sheet_name="Lesion_ge4",
            index=False
        )

        lesion_ge3.to_excel(
            writer,
            sheet_name="Lesion_ge3",
            index=False
        )

        patient_ordinal_auc.to_excel(
            writer,
            sheet_name="Patient_AUC",
            index=False
        )

        patient_ge4.to_excel(
            writer,
            sheet_name="Patient_ge4",
            index=False
        )

        patient_ge3.to_excel(
            writer,
            sheet_name="Patient_ge3",
            index=False
        )

        # pairwise

        lesion_pair_ordinal.to_excel(
            writer,
            sheet_name="L_pair_AUC",
            index=False
        )

        lesion_pair_ge4.to_excel(
            writer,
            sheet_name="L_pair_ge4",
            index=False
        )

        lesion_pair_ge3.to_excel(
            writer,
            sheet_name="L_pair_ge3",
            index=False
        )

        patient_pair_ordinal.to_excel(
            writer,
            sheet_name="P_pair_AUC",
            index=False
        )

        patient_pair_ge4.to_excel(
            writer,
            sheet_name="P_pair_ge4",
            index=False
        )

        patient_pair_ge3.to_excel(
            writer,
            sheet_name="P_pair_ge3",
            index=False
        )

    format_excel(
        excel_file
    )

    # ========================================================
    # CSV output
    # ========================================================

    output_tables = {

        "Lesion_level_AUC.csv":
            lesion_ordinal_auc,

        "Lesion_level_VIRADS_ge4.csv":
            lesion_ge4,

        "Lesion_level_VIRADS_ge3.csv":
            lesion_ge3,

        "Patient_level_AUC.csv":
            patient_ordinal_auc,

        "Patient_level_VIRADS_ge4.csv":
            patient_ge4,

        "Patient_level_VIRADS_ge3.csv":
            patient_ge3,

        "Lesion_pairwise_ordinal_AUC.csv":
            lesion_pair_ordinal,

        "Lesion_pairwise_ge4_AUC.csv":
            lesion_pair_ge4,

        "Lesion_pairwise_ge3_AUC.csv":
            lesion_pair_ge3,

        "Patient_pairwise_ordinal_AUC.csv":
            patient_pair_ordinal,

        "Patient_pairwise_ge4_AUC.csv":
            patient_pair_ge4,

        "Patient_pairwise_ge3_AUC.csv":
            patient_pair_ge3
    }

    for filename, table in output_tables.items():

        table.to_csv(
            os.path.join(
                output_dir,
                filename
            ),
            index=False,
            encoding="utf-8-sig"
        )

    # ========================================================
    # Console summary
    # ========================================================

    print("\n" + "=" * 72)

    print(
        "PAIRWISE ANALYSIS SUMMARY"
    )

    print("=" * 72)

    summary_columns = [
        "Dataset",
        "Better model",
        "Worse model",
        "AUC better",
        "AUC worse",
        "Delta AUC",
        "Delta 95% CI lower",
        "Delta 95% CI upper",
        "P raw",
        "P Holm"
    ]

    for title, table in [

        (
            "LESION ORDINAL",
            lesion_pair_ordinal
        ),

        (
            "LESION >=4",
            lesion_pair_ge4
        ),

        (
            "LESION >=3",
            lesion_pair_ge3
        ),

        (
            "PATIENT ORDINAL",
            patient_pair_ordinal
        ),

        (
            "PATIENT >=4",
            patient_pair_ge4
        ),

        (
            "PATIENT >=3",
            patient_pair_ge3
        )

    ]:

        print("\n" + "-" * 72)

        print(
            title
        )

        print("-" * 72)

        if len(
            table
        ) == 0:

            print(
                "No results."
            )

        else:

            print(
                table[
                    summary_columns
                ].to_string(
                    index=False
                )
            )

    # ========================================================
    # Done
    # ========================================================

    print("\n" + "=" * 72)

    print(
        "分析完成"
    )

    print("=" * 72)

    print(
        f"\n结果目录：\n{output_dir}"
    )

    print(
        f"\nExcel：\n{excel_file}"
    )

    print(
        "\n主要 pairwise sheets:"
    )

    print(
        "L_pair_AUC"
    )

    print(
        "L_pair_ge4"
    )

    print(
        "L_pair_ge3"
    )

    print(
        "P_pair_AUC"
    )

    print(
        "P_pair_ge4"
    )

    print(
        "P_pair_ge3"
    )


# ============================================================
# 23. Run
# ============================================================

if __name__ == "__main__":

    try:

        main()

    except Exception as e:

        print("\n" + "=" * 72)

        print(
            "程序发生错误"
        )

        print("=" * 72)

        print(
            f"\n{type(e).__name__}: {e}"
        )

        print(
            "\n请检查："
        )

        print(
            "1. Group 是否为 0/1"
        )

        print(
            "2. PatientID 是否存在"
        )

        print(
            "3. lesion_id 是否存在"
        )

        print(
            "4. cohort 是否存在"
        )

        print(
            "5. 7个score列名是否完全一致"
        )

        print(
            "6. VI-RADS 是否均为1–5"
        )

        print(
            "7. 每个cohort是否同时有Group=0和Group=1"
        )

        print(
            "8. 7个模型是否存在缺失score"
        )

    finally:

        try:

            input(
                "\n按 Enter 键退出..."
            )

        except Exception:

            pass
# ============================================================
# Fast Direct VI-RADS ROC Analysis
# Lesion-level + Patient-level
# 7 models, all pairwise comparisons
#
# ============================================================
#
# 分析框架
#
# A. Lesion-level
#
# 每一行 = 1 lesion
#
# Group:
#     lesion-level Group
#
# VI-RADS:
#     lesion-level score
#
# Bootstrap:
#     PatientID 为 cluster
#     同一患者所有 lesions 一起被抽中
#
# 分析：
#     1. Ordinal VI-RADS 1–5 AUC
#     2. VI-RADS >=4 threshold-specific AUC
#     3. VI-RADS >=3 threshold-specific AUC
#
#
# B. Patient-level
#
# 每个 PatientID 聚合成 1 行
#
# Patient Group:
#     max(Group)
#
# Patient VI-RADS:
#     每个评分来源 max(VI-RADS)
#
# Bootstrap:
#     ordinary patient bootstrap
#
# 分析：
#     1. Ordinal VI-RADS 1–5 AUC
#     2. VI-RADS >=4 threshold-specific AUC
#     3. VI-RADS >=3 threshold-specific AUC
#
#
# C. Pairwise comparisons
#
# 7 models -> 21 pairwise comparisons
#
# 对每一种分析：
#
#     Better model = observed AUC 较大的模型
#     Worse model  = observed AUC 较小的模型
#
#     Delta AUC =
#         AUC(Better) - AUC(Worse)
#
# Bootstrap difference:
#
#     Delta_boot =
#         AUC_boot(Better) - AUC_boot(Worse)
#
# 注意：
#     bootstrap 内不重新判断谁大谁小。
#
# 输出：
#
#     Delta AUC
#     bootstrap 95% CI
#     raw two-sided bootstrap P
#     Holm-adjusted P
#
#
# D. Speed optimization
#
# 1. 每个 cohort / level 只生成一次 bootstrap
# 2. 7 个模型共享完全相同 bootstrap samples
# 3. 不再 pair-by-pair bootstrap
# 4. lesion-level 不再反复 pd.concat
# 5. 使用 patient bootstrap counts -> sample weights
# 6. VI-RADS 只有少数离散值，使用快速 weighted AUC
#
#
# 不使用：
#     Logistic regression
#     Youden cutoff
#     DeLong
#
# ============================================================


import os
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from sklearn.metrics import (
    roc_curve,
    confusion_matrix
)

from tkinter import Tk, filedialog


warnings.filterwarnings("ignore")


# ============================================================
# 0. 用户设置
# ============================================================

OUTCOME_COLUMN = "Group"

PATIENT_ID_COLUMN = "PatientID"

COHORT_COLUMN = "cohort"

LESION_ID_COLUMN = "lesion_id"


SCORE_COLUMNS = [
    "Reference",
    "Qwen3",
    "Original",
    "GPT4",
    "Radiologists",
    "MedGemma",
    "Deepseek"
]


DISPLAY_NAMES = [
    "Reference",
    "Qwen3",
    "Original reports",
    "GPT-4",
    "Radiologists",
    "MedGemma",
    "DeepSeek"
]


COLORS = [
    "#222222",
    "#D55E00",
    "#999999",
    "#CC79A7",
    "#0072B2",
    "#009E73",
    "#E69F00"
]


# ------------------------------------------------------------
# Bootstrap次数
#
# 2000 与你原来的代码一致。
#
# 因为现在共享 bootstrap + 快速 AUC，
# 会远快于逐 pair bootstrap。
# ------------------------------------------------------------

N_BOOTSTRAP = 2000

RANDOM_SEED = 42

DPI = 600


# ------------------------------------------------------------
# Pairwise比较要求完整scores
#
# 推荐保持 True。
#
# 这样7个模型一定基于完全相同病例进行比较，
# paired comparison 最干净。
#
# 如果某个score确实有缺失，会直接报错提醒。
# ------------------------------------------------------------

REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE = True


# ============================================================
# 1. Configuration check
# ============================================================

if len(SCORE_COLUMNS) != len(DISPLAY_NAMES):

    raise ValueError(
        "SCORE_COLUMNS 和 DISPLAY_NAMES 数量必须一致。"
    )


if len(SCORE_COLUMNS) != len(COLORS):

    raise ValueError(
        "SCORE_COLUMNS 和 COLORS 数量必须一致。"
    )


N_MODELS = len(SCORE_COLUMNS)


# ============================================================
# 2. Load data
# ============================================================

def load_data(file_path):

    file_lower = file_path.lower()

    if file_lower.endswith(".xlsx"):

        df = pd.read_excel(
            file_path,
            engine="openpyxl"
        )

    elif file_lower.endswith(".xls"):

        df = pd.read_excel(
            file_path
        )

    elif file_lower.endswith(".csv"):

        try:

            df = pd.read_csv(
                file_path,
                encoding="utf-8-sig"
            )

        except UnicodeDecodeError:

            df = pd.read_csv(
                file_path,
                encoding="gbk"
            )

    else:

        raise ValueError(
            "仅支持 .xlsx / .xls / .csv 文件。"
        )

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    return df


# ============================================================
# 3. Safe divide
# ============================================================

def safe_divide(a, b):

    if b == 0:

        return np.nan

    return a / b


# ============================================================
# 4. Fast weighted AUC
#
# 与 roc_auc_score 对离散 ordinal / binary predictor
# 的定义一致：
#
# AUC =
# P(score_positive > score_negative)
# +
# 0.5 * P(score_positive == score_negative)
#
# 支持 sample weights。
#
# 对 cluster bootstrap：
# 一个患者被抽中 k 次，
# 该患者所有 lesion 的 weight = k。
#
# 避免真正复制 DataFrame。
# ============================================================

def fast_weighted_auc(
    y_true,
    y_score,
    weights=None
):

    y_true = np.asarray(
        y_true,
        dtype=float
    )

    y_score = np.asarray(
        y_score,
        dtype=float
    )

    if weights is None:

        weights = np.ones(
            len(y_true),
            dtype=float
        )

    else:

        weights = np.asarray(
            weights,
            dtype=float
        )

    valid = (
        np.isfinite(y_true)
        &
        np.isfinite(y_score)
        &
        np.isfinite(weights)
        &
        (weights > 0)
    )

    if not np.any(valid):

        return np.nan

    y = y_true[valid].astype(int)

    score = y_score[valid]

    w = weights[valid]

    positive_weight = w[
        y == 1
    ].sum()

    negative_weight = w[
        y == 0
    ].sum()

    if (
        positive_weight <= 0
        or negative_weight <= 0
    ):

        return np.nan

    score_levels = np.sort(
        np.unique(score)
    )

    cumulative_negative = 0.0

    favorable_pairs = 0.0

    for level in score_levels:

        mask_level = (
            score == level
        )

        pos_at_level = w[
            mask_level & (y == 1)
        ].sum()

        neg_at_level = w[
            mask_level & (y == 0)
        ].sum()

        favorable_pairs += (
            pos_at_level
            *
            (
                cumulative_negative
                +
                0.5 * neg_at_level
            )
        )

        cumulative_negative += (
            neg_at_level
        )

    auc = (
        favorable_pairs
        /
        (
            positive_weight
            *
            negative_weight
        )
    )

    return float(auc)


# ============================================================
# 5. Diagnostic metrics
# ============================================================

def diagnostic_metrics(
    y_true,
    score,
    threshold
):

    y_true = np.asarray(
        y_true,
        dtype=int
    )

    score = np.asarray(
        score,
        dtype=float
    )

    valid = (
        np.isfinite(y_true)
        &
        np.isfinite(score)
    )

    y_true = y_true[
        valid
    ]

    score = score[
        valid
    ]

    y_pred = (
        score >= threshold
    ).astype(int)

    cm = confusion_matrix(
        y_true,
        y_pred,
        labels=[
            0,
            1
        ]
    )

    tn, fp, fn, tp = cm.ravel()

    sensitivity = safe_divide(
        tp,
        tp + fn
    )

    specificity = safe_divide(
        tn,
        tn + fp
    )

    ppv = safe_divide(
        tp,
        tp + fp
    )

    npv = safe_divide(
        tn,
        tn + fn
    )

    accuracy = safe_divide(
        tp + tn,
        tp + tn + fp + fn
    )

    balanced_accuracy = (
        sensitivity + specificity
    ) / 2

    return {
        "TP": int(tp),
        "TN": int(tn),
        "FP": int(fp),
        "FN": int(fn),
        "Sensitivity": sensitivity,
        "Specificity": specificity,
        "PPV": ppv,
        "NPV": npv,
        "Accuracy": accuracy,
        "Balanced Accuracy": balanced_accuracy
    }


# ============================================================
# 6. Data check
# ============================================================

def check_dataframe(
    df,
    dataset_name
):

    print("\n" + "=" * 72)

    print(
        f"检查 cohort：{dataset_name}"
    )

    print("=" * 72)

    print(
        f"Lesion 数量：{len(df)}"
    )

    required_columns = [
        OUTCOME_COLUMN,
        PATIENT_ID_COLUMN,
        LESION_ID_COLUMN
    ] + SCORE_COLUMNS

    missing_columns = [
        col
        for col in required_columns
        if col not in df.columns
    ]

    if missing_columns:

        raise ValueError(
            f"\n{dataset_name} 缺少以下列：\n"
            f"{missing_columns}"
        )

    n_patients = (
        df[PATIENT_ID_COLUMN]
        .dropna()
        .nunique()
    )

    print(
        f"Patient 数量：{n_patients}"
    )

    y = pd.to_numeric(
        df[OUTCOME_COLUMN],
        errors="coerce"
    )

    unique_y = sorted(
        y.dropna()
        .unique()
        .tolist()
    )

    print(
        f"{OUTCOME_COLUMN} 唯一值：{unique_y}"
    )

    if not set(unique_y).issubset(
        {0, 1}
    ):

        raise ValueError(
            f"{OUTCOME_COLUMN} 必须为 0/1。"
        )

    if len(unique_y) < 2:

        raise ValueError(
            f"{dataset_name} 中没有同时包含 0 和 1。"
        )

    print(
        "\nLesion-level Group distribution:"
    )

    print(
        y.value_counts(
            dropna=False
        ).sort_index()
    )

    print(
        "\nVI-RADS score ranges:"
    )

    score_missing = {}

    for score_col in SCORE_COLUMNS:

        values = pd.to_numeric(
            df[score_col],
            errors="coerce"
        )

        valid = values.dropna()

        score_missing[
            score_col
        ] = int(
            values.isna().sum()
        )

        if len(valid) == 0:

            print(
                f"{score_col:<20}: 全部缺失"
            )

            continue

        print(
            f"{score_col:<20}: "
            f"min={valid.min():.0f}, "
            f"max={valid.max():.0f}, "
            f"N={len(valid)}, "
            f"missing={values.isna().sum()}"
        )

        invalid = valid[
            ~valid.isin(
                [1, 2, 3, 4, 5]
            )
        ]

        if len(invalid) > 0:

            raise ValueError(
                f"{score_col} 中存在非1–5评分："
                f"{sorted(invalid.unique())}"
            )

    if REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE:

        incomplete = {
            k: v
            for k, v in score_missing.items()
            if v > 0
        }

        if len(incomplete) > 0:

            raise ValueError(
                "\nPairwise paired comparison 要求7个score"
                "在当前cohort中均完整。\n"
                f"发现缺失：{incomplete}\n"
                "请先检查数据，或将 "
                "REQUIRE_COMPLETE_SCORES_FOR_PAIRWISE=False。"
            )


# ============================================================
# 7. Build patient-level data
# ============================================================

def build_patient_level_data(
    df,
    dataset_name=""
):

    required_columns = [
        PATIENT_ID_COLUMN,
        OUTCOME_COLUMN
    ] + SCORE_COLUMNS

    temp = df[
        required_columns
    ].copy()

    temp[OUTCOME_COLUMN] = pd.to_numeric(
        temp[OUTCOME_COLUMN],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        temp[score_col] = pd.to_numeric(
            temp[score_col],
            errors="coerce"
        )

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    )

    if len(temp) == 0:

        return pd.DataFrame(
            columns=required_columns
        )

    group_nunique = (
        temp
        .groupby(
            PATIENT_ID_COLUMN
        )[OUTCOME_COLUMN]
        .nunique()
    )

    discordant_patients = (
        group_nunique[
            group_nunique > 1
        ]
        .index
        .tolist()
    )

    print("\n" + "-" * 72)

    print(
        f"{dataset_name}: 构建 patient-level 数据"
    )

    print("-" * 72)

    if len(discordant_patients) > 0:

        print(
            f"发现 {len(discordant_patients)} 名患者"
            "同时具有 Group=0 和 Group=1 lesions。"
        )

        print(
            "Patient-level Group = max(Group)。"
        )

    aggregation_dict = {
        OUTCOME_COLUMN: "max"
    }

    for score_col in SCORE_COLUMNS:

        aggregation_dict[
            score_col
        ] = "max"

    patient_df = (
        temp
        .groupby(
            PATIENT_ID_COLUMN,
            as_index=False
        )
        .agg(
            aggregation_dict
        )
    )

    patient_df[
        OUTCOME_COLUMN
    ] = patient_df[
        OUTCOME_COLUMN
    ].astype(int)

    print(
        f"Patient 数量：{len(patient_df)}"
    )

    print(
        "Patient-level pathology distribution:"
    )

    print(
        patient_df[
            OUTCOME_COLUMN
        ]
        .value_counts()
        .sort_index()
    )

    return patient_df


# ============================================================
# 8. Bootstrap matrix
#
# 返回：
#
# point_auc:
#     shape = (3, 7)
#
# bootstrap_auc:
#     shape = (N_BOOTSTRAP, 3, 7)
#
# analysis index:
#
#     0 = Ordinal
#     1 = >=4
#     2 = >=3
#
# ============================================================

ANALYSIS_NAMES = [
    "Ordinal",
    ">=4",
    ">=3"
]


def transform_score(
    score,
    analysis_index
):

    if analysis_index == 0:

        return score

    elif analysis_index == 1:

        return (
            score >= 4
        ).astype(float)

    elif analysis_index == 2:

        return (
            score >= 3
        ).astype(float)

    else:

        raise ValueError(
            "Unknown analysis index."
        )


# ============================================================
# 9. Lesion-level bootstrap matrix
#
# Bootstrap cluster = PatientID
#
# Multinomial counts 与：
#
# np.random.choice(
#     patients,
#     size=n_patients,
#     replace=True
# )
#
# 完全等价。
#
# 若 patient 被抽中 k 次：
# 其所有 lesion weight = k
# ============================================================

def bootstrap_matrix_lesion(
    df,
    n_bootstrap=N_BOOTSTRAP,
    seed=RANDOM_SEED
):

    temp = df[
        [
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ] + SCORE_COLUMNS
    ].copy()

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).reset_index(
        drop=True
    )

    y = temp[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    scores = (
        temp[
            SCORE_COLUMNS
        ]
        .astype(float)
        .to_numpy()
    )

    patient_codes, unique_patients = pd.factorize(
        temp[
            PATIENT_ID_COLUMN
        ],
        sort=False
    )

    n_patients = len(
        unique_patients
    )

    n_rows = len(temp)

    print(
        f"\nLesion bootstrap: "
        f"{n_rows} lesions / "
        f"{n_patients} patients / "
        f"{n_bootstrap} replicates"
    )

    point_auc = np.full(
        (
            3,
            N_MODELS
        ),
        np.nan
    )

    bootstrap_auc = np.full(
        (
            n_bootstrap,
            3,
            N_MODELS
        ),
        np.nan,
        dtype=float
    )

    # --------------------------------------------------------
    # observed AUC
    # --------------------------------------------------------

    for analysis_idx in range(3):

        for model_idx in range(
            N_MODELS
        ):

            predictor = transform_score(
                scores[
                    :,
                    model_idx
                ],
                analysis_idx
            )

            point_auc[
                analysis_idx,
                model_idx
            ] = fast_weighted_auc(
                y,
                predictor
            )

    # --------------------------------------------------------
    # bootstrap
    # --------------------------------------------------------

    rng = np.random.default_rng(
        seed
    )

    probabilities = np.full(
        n_patients,
        1.0 / n_patients
    )

    progress_marks = set(
        [
            int(
                n_bootstrap * x
            )
            for x in [
                0.25,
                0.50,
                0.75,
                1.00
            ]
        ]
    )

    for b in range(
        n_bootstrap
    ):

        patient_counts = rng.multinomial(
            n_patients,
            probabilities
        )

        row_weights = patient_counts[
            patient_codes
        ].astype(float)

        for model_idx in range(
            N_MODELS
        ):

            score_original = scores[
                :,
                model_idx
            ]

            # ordinal

            bootstrap_auc[
                b,
                0,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_original,
                row_weights
            )

            # >=4

            score_ge4 = (
                score_original >= 4
            ).astype(float)

            bootstrap_auc[
                b,
                1,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge4,
                row_weights
            )

            # >=3

            score_ge3 = (
                score_original >= 3
            ).astype(float)

            bootstrap_auc[
                b,
                2,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge3,
                row_weights
            )

        current = b + 1

        if current in progress_marks:

            pct = (
                100 * current / n_bootstrap
            )

            print(
                f"  Lesion bootstrap: "
                f"{pct:.0f}%"
            )

    return (
        point_auc,
        bootstrap_auc
    )


# ============================================================
# 10. Patient-level bootstrap matrix
# ============================================================

def bootstrap_matrix_patient(
    patient_df,
    n_bootstrap=N_BOOTSTRAP,
    seed=RANDOM_SEED
):

    temp = patient_df[
        [
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ] + SCORE_COLUMNS
    ].copy()

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).reset_index(
        drop=True
    )

    y = temp[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    scores = (
        temp[
            SCORE_COLUMNS
        ]
        .astype(float)
        .to_numpy()
    )

    n_patients = len(temp)

    print(
        f"\nPatient bootstrap: "
        f"{n_patients} patients / "
        f"{n_bootstrap} replicates"
    )

    point_auc = np.full(
        (
            3,
            N_MODELS
        ),
        np.nan
    )

    bootstrap_auc = np.full(
        (
            n_bootstrap,
            3,
            N_MODELS
        ),
        np.nan,
        dtype=float
    )

    # --------------------------------------------------------
    # observed
    # --------------------------------------------------------

    for analysis_idx in range(3):

        for model_idx in range(
            N_MODELS
        ):

            predictor = transform_score(
                scores[
                    :,
                    model_idx
                ],
                analysis_idx
            )

            point_auc[
                analysis_idx,
                model_idx
            ] = fast_weighted_auc(
                y,
                predictor
            )

    # --------------------------------------------------------
    # bootstrap
    # --------------------------------------------------------

    rng = np.random.default_rng(
        seed
    )

    probabilities = np.full(
        n_patients,
        1.0 / n_patients
    )

    progress_marks = set(
        [
            int(
                n_bootstrap * x
            )
            for x in [
                0.25,
                0.50,
                0.75,
                1.00
            ]
        ]
    )

    for b in range(
        n_bootstrap
    ):

        patient_weights = rng.multinomial(
            n_patients,
            probabilities
        ).astype(float)

        for model_idx in range(
            N_MODELS
        ):

            score_original = scores[
                :,
                model_idx
            ]

            bootstrap_auc[
                b,
                0,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_original,
                patient_weights
            )

            score_ge4 = (
                score_original >= 4
            ).astype(float)

            bootstrap_auc[
                b,
                1,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge4,
                patient_weights
            )

            score_ge3 = (
                score_original >= 3
            ).astype(float)

            bootstrap_auc[
                b,
                2,
                model_idx
            ] = fast_weighted_auc(
                y,
                score_ge3,
                patient_weights
            )

        current = b + 1

        if current in progress_marks:

            pct = (
                100 * current / n_bootstrap
            )

            print(
                f"  Patient bootstrap: "
                f"{pct:.0f}%"
            )

    return (
        point_auc,
        bootstrap_auc
    )


# ============================================================
# 11. Bootstrap CI
# ============================================================

def bootstrap_ci(
    bootstrap_values
):

    values = np.asarray(
        bootstrap_values,
        dtype=float
    )

    values = values[
        np.isfinite(values)
    ]

    if len(values) == 0:

        return (
            np.nan,
            np.nan
        )

    lower = np.percentile(
        values,
        2.5
    )

    upper = np.percentile(
        values,
        97.5
    )

    return (
        float(lower),
        float(upper)
    )


# ============================================================
# 12. Two-sided paired bootstrap P
#
# observed better-worse orientation固定。
#
# P =
# 2 * min[
#     Pr(Delta_boot <= 0),
#     Pr(Delta_boot >= 0)
# ]
#
# +1 correction 防止出现P=0。
# ============================================================

def paired_bootstrap_pvalue(
    delta_boot
):

    delta_boot = np.asarray(
        delta_boot,
        dtype=float
    )

    delta_boot = delta_boot[
        np.isfinite(delta_boot)
    ]

    n = len(
        delta_boot
    )

    if n == 0:

        return np.nan

    n_le_zero = np.sum(
        delta_boot <= 0
    )

    n_ge_zero = np.sum(
        delta_boot >= 0
    )

    p_lower = (
        n_le_zero + 1
    ) / (
        n + 1
    )

    p_upper = (
        n_ge_zero + 1
    ) / (
        n + 1
    )

    p = 2 * min(
        p_lower,
        p_upper
    )

    return min(
        float(p),
        1.0
    )


# ============================================================
# 13. Holm correction
# ============================================================

def holm_adjust(
    p_values
):

    p_values = np.asarray(
        p_values,
        dtype=float
    )

    adjusted = np.full(
        len(p_values),
        np.nan
    )

    valid_idx = np.where(
        np.isfinite(
            p_values
        )
    )[0]

    if len(valid_idx) == 0:

        return adjusted

    valid_p = p_values[
        valid_idx
    ]

    order = np.argsort(
        valid_p
    )

    sorted_p = valid_p[
        order
    ]

    m = len(
        sorted_p
    )

    adjusted_sorted = np.zeros(
        m,
        dtype=float
    )

    running_max = 0.0

    for i in range(
        m
    ):

        value = (
            (m - i)
            *
            sorted_p[i]
        )

        value = min(
            value,
            1.0
        )

        running_max = max(
            running_max,
            value
        )

        adjusted_sorted[i] = (
            running_max
        )

    reverse_order = np.empty(
        m,
        dtype=int
    )

    reverse_order[
        order
    ] = np.arange(
        m
    )

    adjusted_valid = adjusted_sorted[
        reverse_order
    ]

    adjusted[
        valid_idx
    ] = adjusted_valid

    return adjusted


# ============================================================
# 14. Build individual AUC table
# ============================================================

def build_auc_table(
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc,
    n_lesions=None,
    n_patients=None
):

    rows = []

    if level_name == "Lesion":

        ci_method = (
            "Patient-level cluster bootstrap"
        )

    else:

        ci_method = (
            "Patient-level bootstrap"
        )

    for analysis_idx, analysis_name in enumerate(
        ANALYSIS_NAMES
    ):

        for model_idx, score_col in enumerate(
            SCORE_COLUMNS
        ):

            auc_value = point_auc[
                analysis_idx,
                model_idx
            ]

            lower, upper = bootstrap_ci(
                bootstrap_auc[
                    :,
                    analysis_idx,
                    model_idx
                ]
            )

            row = {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    level_name,

                "Analysis":
                    analysis_name,

                "Variable":
                    score_col,

                "Model":
                    DISPLAY_NAMES[
                        model_idx
                    ],

                "AUC":
                    auc_value,

                "95% CI lower":
                    lower,

                "95% CI upper":
                    upper,

                "AUC (95% CI)":
                    (
                        f"{auc_value:.3f} "
                        f"({lower:.3f}–{upper:.3f})"
                    ),

                "CI method":
                    ci_method,

                "Bootstrap N":
                    N_BOOTSTRAP
            }

            if n_lesions is not None:

                row[
                    "Lesions"
                ] = n_lesions

            if n_patients is not None:

                row[
                    "Patients"
                ] = n_patients

            rows.append(
                row
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# 15. Build pairwise comparison table
# ============================================================

def build_pairwise_table(
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc
):

    all_rows = []

    for analysis_idx, analysis_name in enumerate(
        ANALYSIS_NAMES
    ):

        current_rows = []

        for i in range(
            N_MODELS
        ):

            for j in range(
                i + 1,
                N_MODELS
            ):

                auc_i = point_auc[
                    analysis_idx,
                    i
                ]

                auc_j = point_auc[
                    analysis_idx,
                    j
                ]

                if (
                    not np.isfinite(
                        auc_i
                    )
                    or
                    not np.isfinite(
                        auc_j
                    )
                ):

                    continue

                # --------------------------------------------
                # observed AUC determines direction
                # --------------------------------------------

                if auc_i >= auc_j:

                    better_idx = i

                    worse_idx = j

                else:

                    better_idx = j

                    worse_idx = i

                auc_better = point_auc[
                    analysis_idx,
                    better_idx
                ]

                auc_worse = point_auc[
                    analysis_idx,
                    worse_idx
                ]

                delta_auc = (
                    auc_better
                    -
                    auc_worse
                )

                delta_boot = (
                    bootstrap_auc[
                        :,
                        analysis_idx,
                        better_idx
                    ]
                    -
                    bootstrap_auc[
                        :,
                        analysis_idx,
                        worse_idx
                    ]
                )

                valid_delta = delta_boot[
                    np.isfinite(
                        delta_boot
                    )
                ]

                delta_lower, delta_upper = (
                    bootstrap_ci(
                        valid_delta
                    )
                )

                p_raw = paired_bootstrap_pvalue(
                    valid_delta
                )

                current_rows.append(
                    {
                        "Dataset":
                            dataset_name,

                        "Analysis level":
                            level_name,

                        "Analysis":
                            analysis_name,

                        "Model 1":
                            DISPLAY_NAMES[i],

                        "Model 2":
                            DISPLAY_NAMES[j],

                        "Better model":
                            DISPLAY_NAMES[
                                better_idx
                            ],

                        "Worse model":
                            DISPLAY_NAMES[
                                worse_idx
                            ],

                        "AUC better":
                            auc_better,

                        "AUC worse":
                            auc_worse,

                        "Delta AUC":
                            delta_auc,

                        "Delta 95% CI lower":
                            delta_lower,

                        "Delta 95% CI upper":
                            delta_upper,

                        "Delta AUC (95% CI)":
                            (
                                f"{delta_auc:.3f} "
                                f"({delta_lower:.3f}–"
                                f"{delta_upper:.3f})"
                            ),

                        "P raw":
                            p_raw,

                        "Valid bootstrap replicates":
                            len(
                                valid_delta
                            ),

                        "Comparison method":
                            (
                                "Paired bootstrap; "
                                "shared bootstrap samples"
                            )
                    }
                )

        # ----------------------------------------------------
        # Holm correction:
        # 每个 cohort × level × analysis
        # 的21组comparison作为一个family
        # ----------------------------------------------------

        if len(
            current_rows
        ) > 0:

            raw_ps = [
                row[
                    "P raw"
                ]
                for row in current_rows
            ]

            adjusted_ps = holm_adjust(
                raw_ps
            )

            for row, p_adj in zip(
                current_rows,
                adjusted_ps
            ):

                row[
                    "P Holm"
                ] = p_adj

                row[
                    "Significant raw P<0.05"
                ] = (
                    pd.notna(
                        row[
                            "P raw"
                        ]
                    )
                    and
                    row[
                        "P raw"
                    ] < 0.05
                )

                row[
                    "Significant Holm P<0.05"
                ] = (
                    pd.notna(
                        p_adj
                    )
                    and
                    p_adj < 0.05
                )

            all_rows.extend(
                current_rows
            )

    return pd.DataFrame(
        all_rows
    )


# ============================================================
# 16. Threshold diagnostic table
# ============================================================

def build_threshold_metrics_table(
    source_df,
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc
):

    rows = []

    if level_name == "Lesion":

        n_lesions = len(
            source_df
        )

        n_patients = (
            source_df[
                PATIENT_ID_COLUMN
            ]
            .nunique()
        )

        ci_method = (
            "Patient-level cluster bootstrap"
        )

    else:

        n_lesions = None

        n_patients = len(
            source_df
        )

        ci_method = (
            "Patient-level bootstrap"
        )

    y = source_df[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    for model_idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        score = source_df[
            score_col
        ].astype(float).to_numpy()

        for threshold, analysis_idx in [
            (
                4,
                1
            ),
            (
                3,
                2
            )
        ]:

            metrics = diagnostic_metrics(
                y,
                score,
                threshold
            )

            auc_value = point_auc[
                analysis_idx,
                model_idx
            ]

            lower, upper = bootstrap_ci(
                bootstrap_auc[
                    :,
                    analysis_idx,
                    model_idx
                ]
            )

            row = {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    level_name,

                "Variable":
                    score_col,

                "Model":
                    DISPLAY_NAMES[
                        model_idx
                    ],

                "Threshold":
                    f">={threshold}",

                "Threshold-specific AUC":
                    auc_value,

                "AUC 95% CI lower":
                    lower,

                "AUC 95% CI upper":
                    upper,

                "Threshold AUC (95% CI)":
                    (
                        f"{auc_value:.3f} "
                        f"({lower:.3f}–{upper:.3f})"
                    ),

                "TP":
                    metrics[
                        "TP"
                    ],

                "TN":
                    metrics[
                        "TN"
                    ],

                "FP":
                    metrics[
                        "FP"
                    ],

                "FN":
                    metrics[
                        "FN"
                    ],

                "Sensitivity":
                    metrics[
                        "Sensitivity"
                    ],

                "Specificity":
                    metrics[
                        "Specificity"
                    ],

                "PPV":
                    metrics[
                        "PPV"
                    ],

                "NPV":
                    metrics[
                        "NPV"
                    ],

                "Accuracy":
                    metrics[
                        "Accuracy"
                    ],

                "Balanced Accuracy":
                    metrics[
                        "Balanced Accuracy"
                    ],

                "AUC CI method":
                    ci_method
            }

            if n_lesions is not None:

                row[
                    "Lesions"
                ] = n_lesions

            row[
                "Patients"
            ] = n_patients

            rows.append(
                row
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# 17. Plot style
# ============================================================

def set_plot_style():

    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 11,
            "axes.labelsize": 14,
            "axes.titlesize": 15,
            "axes.linewidth": 1.2,
            "xtick.labelsize": 11,
            "ytick.labelsize": 11,
            "xtick.major.size": 5,
            "ytick.major.size": 5,
            "xtick.major.width": 1.1,
            "ytick.major.width": 1.1,
            "legend.fontsize": 8.5,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none"
        }
    )


set_plot_style()


# ============================================================
# 18. ROC plot
#
# 只画 ordinal VI-RADS 1–5 ROC
# 与你原来的图保持一致
# ============================================================

def plot_ordinal_roc(
    source_df,
    dataset_name,
    level_name,
    point_auc,
    bootstrap_auc,
    output_dir
):

    plt.figure(
        figsize=(
            8,
            6
        )
    )

    y = source_df[
        OUTCOME_COLUMN
    ].astype(int).to_numpy()

    for model_idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        score = source_df[
            score_col
        ].astype(float).to_numpy()

        valid = (
            np.isfinite(y)
            &
            np.isfinite(score)
        )

        y_current = y[
            valid
        ]

        score_current = score[
            valid
        ]

        if np.unique(
            y_current
        ).size < 2:

            continue

        fpr, tpr, _ = roc_curve(
            y_current,
            score_current
        )

        auc_value = point_auc[
            0,
            model_idx
        ]

        lower, upper = bootstrap_ci(
            bootstrap_auc[
                :,
                0,
                model_idx
            ]
        )

        label_text = (
            f"{DISPLAY_NAMES[model_idx]} "
            f"(AUC = {auc_value:.3f} "
            f"[{lower:.3f}-{upper:.3f}])"
        )

        plt.plot(
            fpr,
            tpr,
            color=COLORS[
                model_idx
            ],
            linewidth=2,
            label=label_text
        )

    plt.plot(
        [
            0,
            1
        ],
        [
            0,
            1
        ],
        color="navy",
        linestyle="--",
        linewidth=1.2
    )

    plt.xlim(
        [
            -0.04,
            1.04
        ]
    )

    plt.ylim(
        [
            -0.04,
            1.04
        ]
    )

    plt.xlabel(
        "1 - Specificity"
    )

    plt.ylabel(
        "Sensitivity"
    )

    plt.title(
        f"{level_name}-level ROC - "
        f"{dataset_name}"
    )

    plt.legend(
        loc="lower right"
    )

    plt.grid(
        False
    )

    safe_name = (
        str(
            dataset_name
        )
        .replace(
            " ",
            "_"
        )
        .replace(
            "/",
            "_"
        )
        .replace(
            "\\",
            "_"
        )
        .replace(
            ":",
            "_"
        )
    )

    for ext in [
        "pdf",
        "svg",
        "png"
    ]:

        file_name = os.path.join(
            output_dir,
            (
                f"{safe_name}-"
                f"{level_name}-level-ROC."
                f"{ext}"
            )
        )

        if ext == "png":

            plt.savefig(
                file_name,
                dpi=DPI,
                bbox_inches="tight"
            )

        else:

            plt.savefig(
                file_name,
                format=ext,
                bbox_inches="tight"
            )

    plt.close()


# ============================================================
# 19. Analyze one cohort
# ============================================================

def analyze_cohort(
    df_current,
    dataset_name,
    output_dir
):

    print("\n" + "#" * 72)

    print(
        f"COHORT: {dataset_name}"
    )

    print("#" * 72)

    check_dataframe(
        df_current,
        dataset_name
    )

    # ========================================================
    # Lesion level
    # ========================================================

    print("\n" + "=" * 72)

    print(
        f"LESION-LEVEL: {dataset_name}"
    )

    print("=" * 72)

    (
        lesion_point_auc,
        lesion_boot_auc
    ) = bootstrap_matrix_lesion(
        df_current,
        n_bootstrap=N_BOOTSTRAP,
        seed=RANDOM_SEED
    )

    lesion_auc_table = build_auc_table(
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc,
        n_lesions=len(
            df_current
        ),
        n_patients=df_current[
            PATIENT_ID_COLUMN
        ].nunique()
    )

    lesion_pairwise = build_pairwise_table(
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc
    )

    lesion_threshold = (
        build_threshold_metrics_table(
            source_df=df_current,
            dataset_name=dataset_name,
            level_name="Lesion",
            point_auc=lesion_point_auc,
            bootstrap_auc=lesion_boot_auc
        )
    )

    plot_ordinal_roc(
        source_df=df_current,
        dataset_name=dataset_name,
        level_name="Lesion",
        point_auc=lesion_point_auc,
        bootstrap_auc=lesion_boot_auc,
        output_dir=output_dir
    )

    # ========================================================
    # Patient level
    # ========================================================

    print("\n" + "=" * 72)

    print(
        f"PATIENT-LEVEL: {dataset_name}"
    )

    print("=" * 72)

    patient_df = build_patient_level_data(
        df_current,
        dataset_name
    )

    (
        patient_point_auc,
        patient_boot_auc
    ) = bootstrap_matrix_patient(
        patient_df,
        n_bootstrap=N_BOOTSTRAP,
        seed=RANDOM_SEED + 100000
    )

    patient_auc_table = build_auc_table(
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc,
        n_patients=len(
            patient_df
        )
    )

    patient_pairwise = build_pairwise_table(
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc
    )

    patient_threshold = (
        build_threshold_metrics_table(
            source_df=patient_df,
            dataset_name=dataset_name,
            level_name="Patient",
            point_auc=patient_point_auc,
            bootstrap_auc=patient_boot_auc
        )
    )

    plot_ordinal_roc(
        source_df=patient_df,
        dataset_name=dataset_name,
        level_name="Patient",
        point_auc=patient_point_auc,
        bootstrap_auc=patient_boot_auc,
        output_dir=output_dir
    )

    return {
        "lesion_auc":
            lesion_auc_table,

        "lesion_threshold":
            lesion_threshold,

        "lesion_pairwise":
            lesion_pairwise,

        "patient_auc":
            patient_auc_table,

        "patient_threshold":
            patient_threshold,

        "patient_pairwise":
            patient_pairwise
    }


# ============================================================
# 20. Excel formatting
# ============================================================

def format_excel(
    excel_file
):

    try:

        from openpyxl import (
            load_workbook
        )

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment
        )

        from openpyxl.utils import (
            get_column_letter
        )

        wb = load_workbook(
            excel_file
        )

        decimal_columns = [
            "AUC",
            "95% CI lower",
            "95% CI upper",
            "Threshold-specific AUC",
            "AUC 95% CI lower",
            "AUC 95% CI upper",
            "Sensitivity",
            "Specificity",
            "PPV",
            "NPV",
            "Accuracy",
            "Balanced Accuracy",
            "AUC better",
            "AUC worse",
            "Delta AUC",
            "Delta 95% CI lower",
            "Delta 95% CI upper",
            "P raw",
            "P Holm"
        ]

        for ws in wb.worksheets:

            if ws.max_row < 1:

                continue

            for cell in ws[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor="4F81BD"
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            header_map = {
                cell.value:
                    cell.column
                for cell in ws[1]
            }

            for column_cells in ws.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[
                            0
                        ].column
                    )
                )

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value is None
                        else str(
                            cell.value
                        )
                    )

                    max_length = max(
                        max_length,
                        len(
                            value
                        )
                    )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    42
                )

            ws.freeze_panes = (
                "A2"
            )

            ws.auto_filter.ref = (
                ws.dimensions
            )

            for col_name in decimal_columns:

                if col_name not in header_map:

                    continue

                col_idx = header_map[
                    col_name
                ]

                for row_idx in range(
                    2,
                    ws.max_row + 1
                ):

                    ws.cell(
                        row=row_idx,
                        column=col_idx
                    ).number_format = (
                        "0.000"
                    )

        wb.save(
            excel_file
        )

    except Exception as e:

        print(
            f"Excel 格式化跳过：{e}"
        )


# ============================================================
# 21. Safe concat
# ============================================================

def safe_concat(
    list_of_dfs
):

    valid = [
        x
        for x in list_of_dfs
        if (
            x is not None
            and len(x) > 0
        )
    ]

    if len(valid) == 0:

        return pd.DataFrame()

    return pd.concat(
        valid,
        ignore_index=True
    )


# ============================================================
# 22. Main
# ============================================================

def main():

    root = Tk()

    root.withdraw()

    try:

        root.attributes(
            "-topmost",
            True
        )

    except Exception:

        pass

    print("\n" + "=" * 72)

    print(
        "FAST VI-RADS ROC + ALL PAIRWISE AUC COMPARISONS"
    )

    print("=" * 72)

    print(
        f"\nModels: {N_MODELS}"
    )

    print(
        "Pairwise comparisons per analysis: 21"
    )

    print(
        f"Bootstrap replicates: {N_BOOTSTRAP}"
    )

    print(
        "\nAnalysis:"
    )

    print(
        "1. Lesion ordinal"
    )

    print(
        "2. Lesion >=4"
    )

    print(
        "3. Lesion >=3"
    )

    print(
        "4. Patient ordinal"
    )

    print(
        "5. Patient >=4"
    )

    print(
        "6. Patient >=3"
    )

    file_path = (
        filedialog.askopenfilename(
            title=(
                "选择 VI-RADS 数据文件"
            ),
            filetypes=[
                (
                    "Excel / CSV",
                    "*.xlsx *.xls *.csv"
                ),
                (
                    "All files",
                    "*.*"
                )
            ]
        )
    )

    root.destroy()

    if not file_path:

        print(
            "没有选择文件。"
        )

        return

    df = load_data(
        file_path
    )

    print(
        f"\n读取文件：{file_path}"
    )

    print(
        f"数据维度：{df.shape}"
    )

    # --------------------------------------------------------
    # required columns
    # --------------------------------------------------------

    required = [
        COHORT_COLUMN,
        PATIENT_ID_COLUMN,
        LESION_ID_COLUMN,
        OUTCOME_COLUMN
    ] + SCORE_COLUMNS

    missing = [
        col
        for col in required
        if col not in df.columns
    ]

    if missing:

        raise ValueError(
            f"缺少列：{missing}"
        )

    # --------------------------------------------------------
    # cohort cleanup
    # --------------------------------------------------------

    mask = df[
        COHORT_COLUMN
    ].notna()

    df.loc[
        mask,
        COHORT_COLUMN
    ] = (
        df.loc[
            mask,
            COHORT_COLUMN
        ]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    df.loc[
        df[
            COHORT_COLUMN
        ] == "",
        COHORT_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # PatientID cleanup
    # --------------------------------------------------------

    mask = df[
        PATIENT_ID_COLUMN
    ].notna()

    df.loc[
        mask,
        PATIENT_ID_COLUMN
    ] = (
        df.loc[
            mask,
            PATIENT_ID_COLUMN
        ]
        .astype(str)
        .str.strip()
    )

    df.loc[
        df[
            PATIENT_ID_COLUMN
        ] == "",
        PATIENT_ID_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # lesion_id cleanup
    # --------------------------------------------------------

    mask = df[
        LESION_ID_COLUMN
    ].notna()

    df.loc[
        mask,
        LESION_ID_COLUMN
    ] = (
        df.loc[
            mask,
            LESION_ID_COLUMN
        ]
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # numeric conversion
    # --------------------------------------------------------

    df[
        OUTCOME_COLUMN
    ] = pd.to_numeric(
        df[
            OUTCOME_COLUMN
        ],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        df[
            score_col
        ] = pd.to_numeric(
            df[
                score_col
            ],
            errors="coerce"
        )

    # --------------------------------------------------------
    # Remove invalid patient / outcome
    # --------------------------------------------------------

    before = len(
        df
    )

    df = df.dropna(
        subset=[
            COHORT_COLUMN,
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    ).copy()

    after = len(
        df
    )

    if before != after:

        print(
            f"\n删除 {before-after} 行 "
            "cohort / PatientID / Group 缺失记录。"
        )

    # --------------------------------------------------------
    # cohorts
    # --------------------------------------------------------

    cohort_names = (
        df[
            COHORT_COLUMN
        ]
        .dropna()
        .unique()
        .tolist()
    )

    if len(cohort_names) == 0:

        raise ValueError(
            "cohort列没有有效值。"
        )

    print(
        "\nCohorts:"
    )

    print(
        df[
            COHORT_COLUMN
        ]
        .value_counts()
    )

    # --------------------------------------------------------
    # output dir
    # --------------------------------------------------------

    input_dir = os.path.dirname(
        os.path.abspath(
            file_path
        )
    )

    output_dir = os.path.join(
        input_dir,
        "Fast_VIRADS_Pairwise_ROC_results"
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    all_lesion_auc = []

    all_lesion_threshold = []

    all_lesion_pairwise = []

    all_patient_auc = []

    all_patient_threshold = []

    all_patient_pairwise = []

    # ========================================================
    # cohort loop
    # ========================================================

    for cohort_number, cohort_name in enumerate(
        cohort_names
    ):

        df_current = df[
            df[
                COHORT_COLUMN
            ] == cohort_name
        ].copy()

        results = analyze_cohort(
            df_current=df_current,
            dataset_name=str(
                cohort_name
            ),
            output_dir=output_dir
        )

        all_lesion_auc.append(
            results[
                "lesion_auc"
            ]
        )

        all_lesion_threshold.append(
            results[
                "lesion_threshold"
            ]
        )

        all_lesion_pairwise.append(
            results[
                "lesion_pairwise"
            ]
        )

        all_patient_auc.append(
            results[
                "patient_auc"
            ]
        )

        all_patient_threshold.append(
            results[
                "patient_threshold"
            ]
        )

        all_patient_pairwise.append(
            results[
                "patient_pairwise"
            ]
        )

    # ========================================================
    # combine
    # ========================================================

    lesion_auc = safe_concat(
        all_lesion_auc
    )

    lesion_threshold = safe_concat(
        all_lesion_threshold
    )

    lesion_pairwise = safe_concat(
        all_lesion_pairwise
    )

    patient_auc = safe_concat(
        all_patient_auc
    )

    patient_threshold = safe_concat(
        all_patient_threshold
    )

    patient_pairwise = safe_concat(
        all_patient_pairwise
    )

    # ========================================================
    # Split tables
    # ========================================================

    lesion_ordinal_auc = (
        lesion_auc[
            lesion_auc[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ordinal_auc = (
        patient_auc[
            patient_auc[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_ge4 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_ge3 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge4 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge3 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ordinal = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ge4 = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_pair_ge3 = (
        lesion_pairwise[
            lesion_pairwise[
                "Analysis"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ordinal = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == "Ordinal"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ge4 = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_pair_ge3 = (
        patient_pairwise[
            patient_pairwise[
                "Analysis"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    # ========================================================
    # Excel
    # ========================================================

    excel_file = os.path.join(
        output_dir,
        "VI-RADS_FAST_all_pairwise_analysis.xlsx"
    )

    with pd.ExcelWriter(
        excel_file,
        engine="openpyxl"
    ) as writer:

        # individual performance

        lesion_ordinal_auc.to_excel(
            writer,
            sheet_name="Lesion_AUC",
            index=False
        )

        lesion_ge4.to_excel(
            writer,
            sheet_name="Lesion_ge4",
            index=False
        )

        lesion_ge3.to_excel(
            writer,
            sheet_name="Lesion_ge3",
            index=False
        )

        patient_ordinal_auc.to_excel(
            writer,
            sheet_name="Patient_AUC",
            index=False
        )

        patient_ge4.to_excel(
            writer,
            sheet_name="Patient_ge4",
            index=False
        )

        patient_ge3.to_excel(
            writer,
            sheet_name="Patient_ge3",
            index=False
        )

        # pairwise

        lesion_pair_ordinal.to_excel(
            writer,
            sheet_name="L_pair_AUC",
            index=False
        )

        lesion_pair_ge4.to_excel(
            writer,
            sheet_name="L_pair_ge4",
            index=False
        )

        lesion_pair_ge3.to_excel(
            writer,
            sheet_name="L_pair_ge3",
            index=False
        )

        patient_pair_ordinal.to_excel(
            writer,
            sheet_name="P_pair_AUC",
            index=False
        )

        patient_pair_ge4.to_excel(
            writer,
            sheet_name="P_pair_ge4",
            index=False
        )

        patient_pair_ge3.to_excel(
            writer,
            sheet_name="P_pair_ge3",
            index=False
        )

    format_excel(
        excel_file
    )

    # ========================================================
    # CSV output
    # ========================================================

    output_tables = {

        "Lesion_level_AUC.csv":
            lesion_ordinal_auc,

        "Lesion_level_VIRADS_ge4.csv":
            lesion_ge4,

        "Lesion_level_VIRADS_ge3.csv":
            lesion_ge3,

        "Patient_level_AUC.csv":
            patient_ordinal_auc,

        "Patient_level_VIRADS_ge4.csv":
            patient_ge4,

        "Patient_level_VIRADS_ge3.csv":
            patient_ge3,

        "Lesion_pairwise_ordinal_AUC.csv":
            lesion_pair_ordinal,

        "Lesion_pairwise_ge4_AUC.csv":
            lesion_pair_ge4,

        "Lesion_pairwise_ge3_AUC.csv":
            lesion_pair_ge3,

        "Patient_pairwise_ordinal_AUC.csv":
            patient_pair_ordinal,

        "Patient_pairwise_ge4_AUC.csv":
            patient_pair_ge4,

        "Patient_pairwise_ge3_AUC.csv":
            patient_pair_ge3
    }

    for filename, table in output_tables.items():

        table.to_csv(
            os.path.join(
                output_dir,
                filename
            ),
            index=False,
            encoding="utf-8-sig"
        )

    # ========================================================
    # Console summary
    # ========================================================

    print("\n" + "=" * 72)

    print(
        "PAIRWISE ANALYSIS SUMMARY"
    )

    print("=" * 72)

    summary_columns = [
        "Dataset",
        "Better model",
        "Worse model",
        "AUC better",
        "AUC worse",
        "Delta AUC",
        "Delta 95% CI lower",
        "Delta 95% CI upper",
        "P raw",
        "P Holm"
    ]

    for title, table in [

        (
            "LESION ORDINAL",
            lesion_pair_ordinal
        ),

        (
            "LESION >=4",
            lesion_pair_ge4
        ),

        (
            "LESION >=3",
            lesion_pair_ge3
        ),

        (
            "PATIENT ORDINAL",
            patient_pair_ordinal
        ),

        (
            "PATIENT >=4",
            patient_pair_ge4
        ),

        (
            "PATIENT >=3",
            patient_pair_ge3
        )

    ]:

        print("\n" + "-" * 72)

        print(
            title
        )

        print("-" * 72)

        if len(
            table
        ) == 0:

            print(
                "No results."
            )

        else:

            print(
                table[
                    summary_columns
                ].to_string(
                    index=False
                )
            )

    # ========================================================
    # Done
    # ========================================================

    print("\n" + "=" * 72)

    print(
        "分析完成"
    )

    print("=" * 72)

    print(
        f"\n结果目录：\n{output_dir}"
    )

    print(
        f"\nExcel：\n{excel_file}"
    )

    print(
        "\n主要 pairwise sheets:"
    )

    print(
        "L_pair_AUC"
    )

    print(
        "L_pair_ge4"
    )

    print(
        "L_pair_ge3"
    )

    print(
        "P_pair_AUC"
    )

    print(
        "P_pair_ge4"
    )

    print(
        "P_pair_ge3"
    )


# ============================================================
# 23. Run
# ============================================================

if __name__ == "__main__":

    try:

        main()

    except Exception as e:

        print("\n" + "=" * 72)

        print(
            "程序发生错误"
        )

        print("=" * 72)

        print(
            f"\n{type(e).__name__}: {e}"
        )

        print(
            "\n请检查："
        )

        print(
            "1. Group 是否为 0/1"
        )

        print(
            "2. PatientID 是否存在"
        )

        print(
            "3. lesion_id 是否存在"
        )

        print(
            "4. cohort 是否存在"
        )

        print(
            "5. 7个score列名是否完全一致"
        )

        print(
            "6. VI-RADS 是否均为1–5"
        )

        print(
            "7. 每个cohort是否同时有Group=0和Group=1"
        )

        print(
            "8. 7个模型是否存在缺失score"
        )

    finally:

        try:

            input(
                "\n按 Enter 键退出..."
            )

        except Exception:

            pass
# ============================================================
# Direct VI-RADS ROC analysis
# Lesion-level + Patient-level
# Full ordinal AUC + threshold-specific AUC
# ============================================================
#
# 分析内容：
#
# A. Lesion-level
#    1. VI-RADS 1-5 ordinal ROC / AUC
#       - 95% CI: patient-level cluster bootstrap
#
#    2. VI-RADS >=4
#       - threshold-specific AUC
#       - 95% CI: patient-level cluster bootstrap
#       - Sensitivity / Specificity / PPV / NPV / Accuracy
#       - Balanced Accuracy / TP / TN / FP / FN
#
#    3. VI-RADS >=3
#       - 同上
#
# B. Patient-level
#    Patient outcome:
#       同一患者任意 lesion Group=1
#       -> patient Group=1
#
#    Patient score:
#       每个评分来源取患者所有 lesion 中最高 VI-RADS
#
#    1. VI-RADS 1-5 ordinal ROC / AUC
#       - 95% CI: ordinary patient bootstrap
#
#    2. VI-RADS >=4
#       - threshold-specific AUC
#       - 95% CI: ordinary patient bootstrap
#
#    3. VI-RADS >=3
#       - threshold-specific AUC
#       - 95% CI: ordinary patient bootstrap
#
# 不使用 Logistic Regression
# 不使用 Youden cutoff
#
# ============================================================


import os
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from sklearn.metrics import (
    roc_curve,
    roc_auc_score,
    confusion_matrix
)

from tkinter import Tk, filedialog


warnings.filterwarnings("ignore")


# ============================================================
# 0. 用户设置
# ============================================================

OUTCOME_COLUMN = "Group"

PATIENT_ID_COLUMN = "PatientID"

COHORT_COLUMN = "cohort"

LESION_ID_COLUMN = "lesion_id"


SCORE_COLUMNS = [
    "Reference",
    "Qwen3",
    "Original",
    "GPT4",
    "Radiologists",
    "MedGemma",
    "Deepseek"
]


DISPLAY_NAMES = [
    "Reference",
    "Qwen3",
    "Original reports",
    "GPT-4",
    "Radiologists",
    "MedGemma",
    "DeepSeek"
]


COLORS = [
    "#222222",
    "#D55E00",
    "#999999",
    "#CC79A7",
    "#0072B2",
    "#009E73",
    "#E69F00"
]


N_BOOTSTRAP = 2000

RANDOM_SEED = 42

DPI = 600


# ============================================================
# 1. 配置检查
# ============================================================

if len(SCORE_COLUMNS) != len(DISPLAY_NAMES):

    raise ValueError(
        "SCORE_COLUMNS 和 DISPLAY_NAMES 数量必须一致。"
    )


if len(SCORE_COLUMNS) != len(COLORS):

    raise ValueError(
        "SCORE_COLUMNS 和 COLORS 数量必须一致。"
    )


# ============================================================
# 2. 数据读取
# ============================================================

def load_data(file_path):

    file_lower = file_path.lower()

    if file_lower.endswith(".xlsx"):

        df = pd.read_excel(
            file_path,
            engine="openpyxl"
        )

    elif file_lower.endswith(".xls"):

        df = pd.read_excel(
            file_path
        )

    elif file_lower.endswith(".csv"):

        try:

            df = pd.read_csv(
                file_path,
                encoding="utf-8-sig"
            )

        except UnicodeDecodeError:

            df = pd.read_csv(
                file_path,
                encoding="gbk"
            )

    else:

        raise ValueError(
            "仅支持 .xlsx / .xls / .csv 文件。"
        )

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    return df


# ============================================================
# 3. 安全除法
# ============================================================

def safe_divide(a, b):

    if b == 0:

        return np.nan

    return a / b


# ============================================================
# 4. 数据检查
# ============================================================

def check_dataframe(
    df,
    dataset_name
):

    print("\n" + "=" * 72)

    print(
        f"检查 cohort：{dataset_name}"
    )

    print("=" * 72)

    print(
        f"Lesion 数量：{len(df)}"
    )

    required_columns = [
        OUTCOME_COLUMN,
        PATIENT_ID_COLUMN
    ] + SCORE_COLUMNS

    missing_columns = [
        col
        for col in required_columns
        if col not in df.columns
    ]

    if missing_columns:

        raise ValueError(
            f"\n{dataset_name} 缺少以下列：\n"
            f"{missing_columns}\n\n"
            f"当前列名：\n"
            f"{df.columns.tolist()}"
        )

    n_patients = (
        df[PATIENT_ID_COLUMN]
        .dropna()
        .nunique()
    )

    print(
        f"Patient 数量：{n_patients}"
    )

    y = pd.to_numeric(
        df[OUTCOME_COLUMN],
        errors="coerce"
    )

    unique_y = sorted(
        y.dropna()
        .unique()
        .tolist()
    )

    print(
        f"\n{OUTCOME_COLUMN} 唯一值：{unique_y}"
    )

    if not set(unique_y).issubset(
        {0, 1}
    ):

        raise ValueError(
            f"{OUTCOME_COLUMN} 必须为 0/1。"
        )

    if len(unique_y) < 2:

        raise ValueError(
            f"{dataset_name} 中没有同时包含 0 和 1，"
            "无法计算 ROC。"
        )

    print(
        "\nLesion-level Group 分布："
    )

    print(
        y.value_counts(
            dropna=False
        ).sort_index()
    )

    print(
        "\nVI-RADS 评分范围："
    )

    for score_col in SCORE_COLUMNS:

        values = pd.to_numeric(
            df[score_col],
            errors="coerce"
        )

        valid = values.dropna()

        if len(valid) == 0:

            print(
                f"{score_col:<20}: 全部缺失"
            )

            continue

        print(
            f"{score_col:<20}: "
            f"min={valid.min():.0f}, "
            f"max={valid.max():.0f}, "
            f"N={len(valid)}"
        )

        invalid = valid[
            ~valid.isin(
                [1, 2, 3, 4, 5]
            )
        ]

        if len(invalid) > 0:

            print(
                f"WARNING: {score_col} "
                f"中存在非 1-5 的值："
                f"{sorted(invalid.unique())}"
            )


# ============================================================
# 5. Ordinary bootstrap AUC
#
# Patient-level ordinal AUC 使用
# ============================================================

def bootstrap_auc_standard(
    y_true,
    y_score,
    n_bootstrap=2000,
    seed=42
):

    y_true = np.asarray(
        y_true,
        dtype=int
    )

    y_score = np.asarray(
        y_score,
        dtype=float
    )

    if np.unique(y_true).size < 2:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    original_auc = roc_auc_score(
        y_true,
        y_score
    )

    rng = np.random.default_rng(
        seed
    )

    n = len(y_true)

    bootstrap_aucs = []

    for _ in range(
        n_bootstrap
    ):

        indices = rng.integers(
            0,
            n,
            size=n
        )

        y_boot = y_true[
            indices
        ]

        score_boot = y_score[
            indices
        ]

        if np.unique(y_boot).size < 2:

            continue

        try:

            auc_boot = roc_auc_score(
                y_boot,
                score_boot
            )

            bootstrap_aucs.append(
                auc_boot
            )

        except Exception:

            continue

    if len(bootstrap_aucs) == 0:

        return (
            original_auc,
            np.nan,
            np.nan
        )

    lower = np.percentile(
        bootstrap_aucs,
        2.5
    )

    upper = np.percentile(
        bootstrap_aucs,
        97.5
    )

    return (
        original_auc,
        lower,
        upper
    )


# ============================================================
# 6. Lesion-level cluster bootstrap ordinal AUC
# ============================================================

def bootstrap_auc_cluster(
    y_true,
    y_score,
    patient_ids,
    n_bootstrap=2000,
    seed=42
):

    temp = pd.DataFrame(
        {
            "y": y_true,
            "score": y_score,
            "patient": patient_ids
        }
    )

    temp["y"] = pd.to_numeric(
        temp["y"],
        errors="coerce"
    )

    temp["score"] = pd.to_numeric(
        temp["score"],
        errors="coerce"
    )

    temp = temp.dropna(
        subset=[
            "y",
            "score",
            "patient"
        ]
    ).reset_index(
        drop=True
    )

    if len(temp) == 0:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    temp["y"] = temp["y"].astype(
        int
    )

    temp["score"] = temp["score"].astype(
        float
    )

    if temp["y"].nunique() < 2:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    original_auc = roc_auc_score(
        temp["y"],
        temp["score"]
    )

    unique_patients = (
        temp["patient"]
        .drop_duplicates()
        .to_numpy()
    )

    n_patients = len(
        unique_patients
    )

    rng = np.random.default_rng(
        seed
    )

    bootstrap_aucs = []

    for _ in range(
        n_bootstrap
    ):

        sampled_patients = rng.choice(
            unique_patients,
            size=n_patients,
            replace=True
        )

        boot_parts = []

        for patient in sampled_patients:

            patient_data = temp[
                temp["patient"] == patient
            ]

            boot_parts.append(
                patient_data
            )

        boot_df = pd.concat(
            boot_parts,
            ignore_index=True
        )

        if boot_df["y"].nunique() < 2:

            continue

        try:

            auc_boot = roc_auc_score(
                boot_df["y"],
                boot_df["score"]
            )

            bootstrap_aucs.append(
                auc_boot
            )

        except Exception:

            continue

    if len(bootstrap_aucs) == 0:

        return (
            original_auc,
            np.nan,
            np.nan
        )

    lower = np.percentile(
        bootstrap_aucs,
        2.5
    )

    upper = np.percentile(
        bootstrap_aucs,
        97.5
    )

    return (
        original_auc,
        lower,
        upper
    )


# ============================================================
# 7. Lesion-level threshold-specific AUC
#
# VI-RADS 被二值化：
#
# score >= threshold -> 1
# score < threshold  -> 0
#
# AUC 对 binary predictor 计算。
#
# Bootstrap：
# patient-level cluster bootstrap
# ============================================================

def threshold_auc_cluster_bootstrap(
    y_true,
    score,
    patient_ids,
    threshold,
    n_bootstrap=2000,
    seed=42
):

    temp = pd.DataFrame(
        {
            "y": y_true,
            "score": score,
            "patient": patient_ids
        }
    )

    temp["y"] = pd.to_numeric(
        temp["y"],
        errors="coerce"
    )

    temp["score"] = pd.to_numeric(
        temp["score"],
        errors="coerce"
    )

    temp = temp.dropna(
        subset=[
            "y",
            "score",
            "patient"
        ]
    ).reset_index(
        drop=True
    )

    if len(temp) == 0:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    temp["y"] = temp["y"].astype(
        int
    )

    temp["binary_prediction"] = (
        temp["score"] >= threshold
    ).astype(
        int
    )

    if temp["y"].nunique() < 2:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    original_auc = roc_auc_score(
        temp["y"],
        temp["binary_prediction"]
    )

    unique_patients = (
        temp["patient"]
        .drop_duplicates()
        .to_numpy()
    )

    n_patients = len(
        unique_patients
    )

    rng = np.random.default_rng(
        seed
    )

    bootstrap_aucs = []

    for _ in range(
        n_bootstrap
    ):

        sampled_patients = rng.choice(
            unique_patients,
            size=n_patients,
            replace=True
        )

        boot_parts = []

        for patient in sampled_patients:

            patient_data = temp[
                temp["patient"] == patient
            ]

            boot_parts.append(
                patient_data
            )

        boot_df = pd.concat(
            boot_parts,
            ignore_index=True
        )

        if boot_df["y"].nunique() < 2:

            continue

        try:

            auc_boot = roc_auc_score(
                boot_df["y"],
                boot_df["binary_prediction"]
            )

            bootstrap_aucs.append(
                auc_boot
            )

        except Exception:

            continue

    if len(bootstrap_aucs) == 0:

        return (
            original_auc,
            np.nan,
            np.nan
        )

    lower = np.percentile(
        bootstrap_aucs,
        2.5
    )

    upper = np.percentile(
        bootstrap_aucs,
        97.5
    )

    return (
        original_auc,
        lower,
        upper
    )


# ============================================================
# 8. Patient-level threshold-specific AUC bootstrap
# ============================================================

def threshold_auc_patient_bootstrap(
    y_true,
    score,
    threshold,
    n_bootstrap=2000,
    seed=42
):

    y_true = np.asarray(
        y_true,
        dtype=int
    )

    score = np.asarray(
        score,
        dtype=float
    )

    if np.unique(y_true).size < 2:

        return (
            np.nan,
            np.nan,
            np.nan
        )

    binary_prediction = (
        score >= threshold
    ).astype(
        int
    )

    original_auc = roc_auc_score(
        y_true,
        binary_prediction
    )

    rng = np.random.default_rng(
        seed
    )

    n = len(y_true)

    bootstrap_aucs = []

    for _ in range(
        n_bootstrap
    ):

        indices = rng.integers(
            0,
            n,
            size=n
        )

        y_boot = y_true[
            indices
        ]

        pred_boot = binary_prediction[
            indices
        ]

        if np.unique(y_boot).size < 2:

            continue

        try:

            auc_boot = roc_auc_score(
                y_boot,
                pred_boot
            )

            bootstrap_aucs.append(
                auc_boot
            )

        except Exception:

            continue

    if len(bootstrap_aucs) == 0:

        return (
            original_auc,
            np.nan,
            np.nan
        )

    lower = np.percentile(
        bootstrap_aucs,
        2.5
    )

    upper = np.percentile(
        bootstrap_aucs,
        97.5
    )

    return (
        original_auc,
        lower,
        upper
    )


# ============================================================
# 9. 阈值诊断性能
# ============================================================

def diagnostic_metrics(
    y_true,
    score,
    threshold
):

    y_true = np.asarray(
        y_true,
        dtype=int
    )

    score = np.asarray(
        score,
        dtype=float
    )

    y_pred = (
        score >= threshold
    ).astype(
        int
    )

    cm = confusion_matrix(
        y_true,
        y_pred,
        labels=[
            0,
            1
        ]
    )

    tn, fp, fn, tp = cm.ravel()

    sensitivity = safe_divide(
        tp,
        tp + fn
    )

    specificity = safe_divide(
        tn,
        tn + fp
    )

    ppv = safe_divide(
        tp,
        tp + fp
    )

    npv = safe_divide(
        tn,
        tn + fn
    )

    accuracy = safe_divide(
        tp + tn,
        tp + tn + fp + fn
    )

    balanced_accuracy = (
        sensitivity + specificity
    ) / 2

    return {
        "TP": int(tp),
        "TN": int(tn),
        "FP": int(fp),
        "FN": int(fn),
        "Sensitivity": sensitivity,
        "Specificity": specificity,
        "PPV": ppv,
        "NPV": npv,
        "Accuracy": accuracy,
        "Balanced Accuracy": balanced_accuracy
    }


# ============================================================
# 10. Patient-level 数据
# ============================================================

def build_patient_level_data(
    df,
    dataset_name=""
):

    required_columns = [
        PATIENT_ID_COLUMN,
        OUTCOME_COLUMN
    ] + SCORE_COLUMNS

    temp = df[
        required_columns
    ].copy()

    temp[OUTCOME_COLUMN] = pd.to_numeric(
        temp[OUTCOME_COLUMN],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        temp[score_col] = pd.to_numeric(
            temp[score_col],
            errors="coerce"
        )

    temp = temp.dropna(
        subset=[
            PATIENT_ID_COLUMN,
            OUTCOME_COLUMN
        ]
    )

    if len(temp) == 0:

        return pd.DataFrame(
            columns=required_columns
        )

    group_nunique = (
        temp
        .groupby(
            PATIENT_ID_COLUMN
        )[OUTCOME_COLUMN]
        .nunique()
    )

    discordant_patients = (
        group_nunique[
            group_nunique > 1
        ]
        .index
        .tolist()
    )

    print("\n" + "-" * 72)

    print(
        f"{dataset_name}: 构建 patient-level 数据"
    )

    print("-" * 72)

    if len(discordant_patients) > 0:

        print(
            f"发现 {len(discordant_patients)} 名患者"
            "同时具有 Group=0 与 Group=1 lesions。"
        )

        print(
            "Patient-level Group 按 max(Group) 处理。"
        )

    aggregation_dict = {
        OUTCOME_COLUMN: "max"
    }

    for score_col in SCORE_COLUMNS:

        aggregation_dict[
            score_col
        ] = "max"

    patient_df = (
        temp
        .groupby(
            PATIENT_ID_COLUMN,
            as_index=False
        )
        .agg(
            aggregation_dict
        )
    )

    patient_df[OUTCOME_COLUMN] = (
        patient_df[
            OUTCOME_COLUMN
        ]
        .astype(int)
    )

    print(
        f"\nPatient 数量：{len(patient_df)}"
    )

    print(
        "\nPatient-level pathology distribution:"
    )

    print(
        patient_df[
            OUTCOME_COLUMN
        ]
        .value_counts()
        .sort_index()
    )

    return patient_df


# ============================================================
# 11. 绘图风格
# ============================================================

def set_plot_style():

    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 11,
            "axes.labelsize": 14,
            "axes.titlesize": 15,
            "axes.linewidth": 1.2,
            "xtick.labelsize": 11,
            "ytick.labelsize": 11,
            "xtick.major.size": 5,
            "ytick.major.size": 5,
            "xtick.major.width": 1.1,
            "ytick.major.width": 1.1,
            "legend.fontsize": 8.5,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none"
        }
    )


set_plot_style()


# ============================================================
# 12. Lesion-level analysis
# ============================================================

def analyze_lesion_level_dataset(
    df,
    dataset_name,
    output_dir
):

    print("\n" + "#" * 72)

    print(
        f"LESION-LEVEL：{dataset_name}"
    )

    print("#" * 72)

    check_dataframe(
        df,
        dataset_name
    )

    auc_rows = []

    threshold_rows = []

    roc_objects = {}

    # ========================================================
    # Ordinal 1-5 AUC
    # ========================================================

    for idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        display_name = DISPLAY_NAMES[
            idx
        ]

        temp = df[
            [
                OUTCOME_COLUMN,
                score_col,
                PATIENT_ID_COLUMN
            ]
        ].copy()

        temp[OUTCOME_COLUMN] = pd.to_numeric(
            temp[OUTCOME_COLUMN],
            errors="coerce"
        )

        temp[score_col] = pd.to_numeric(
            temp[score_col],
            errors="coerce"
        )

        temp = temp.dropna(
            subset=[
                OUTCOME_COLUMN,
                score_col,
                PATIENT_ID_COLUMN
            ]
        )

        if len(temp) == 0:

            continue

        y_true = temp[
            OUTCOME_COLUMN
        ].astype(int).to_numpy()

        y_score = temp[
            score_col
        ].astype(float).to_numpy()

        patient_ids = temp[
            PATIENT_ID_COLUMN
        ].to_numpy()

        if np.unique(y_true).size < 2:

            continue

        fpr, tpr, _ = roc_curve(
            y_true,
            y_score
        )

        (
            auc_value,
            lower,
            upper
        ) = bootstrap_auc_cluster(
            y_true=y_true,
            y_score=y_score,
            patient_ids=patient_ids,
            n_bootstrap=N_BOOTSTRAP,
            seed=RANDOM_SEED + idx
        )

        roc_objects[
            score_col
        ] = {
            "fpr": fpr,
            "tpr": tpr,
            "auc": auc_value,
            "lower": lower,
            "upper": upper
        }

        auc_rows.append(
            {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    "Lesion",

                "Variable":
                    score_col,

                "Model":
                    display_name,

                "AUC":
                    auc_value,

                "95% CI lower":
                    lower,

                "95% CI upper":
                    upper,

                "Lesions":
                    len(temp),

                "Patients":
                    temp[PATIENT_ID_COLUMN].nunique(),

                "CI method":
                    "Patient-level cluster bootstrap"
            }
        )

        print(
            f"\n{display_name}: "
            f"AUC={auc_value:.3f} "
            f"({lower:.3f}–{upper:.3f})"
        )

        # ====================================================
        # Threshold-specific analyses
        # ====================================================

        for threshold in [
            4,
            3
        ]:

            metrics = diagnostic_metrics(
                y_true=y_true,
                score=y_score,
                threshold=threshold
            )

            (
                threshold_auc,
                threshold_lower,
                threshold_upper
            ) = threshold_auc_cluster_bootstrap(
                y_true=y_true,
                score=y_score,
                patient_ids=patient_ids,
                threshold=threshold,
                n_bootstrap=N_BOOTSTRAP,
                seed=(
                    RANDOM_SEED
                    + idx * 100
                    + threshold
                )
            )

            threshold_rows.append(
                {
                    "Dataset":
                        dataset_name,

                    "Analysis level":
                        "Lesion",

                    "Variable":
                        score_col,

                    "Model":
                        display_name,

                    "Threshold":
                        f">={threshold}",

                    "Threshold-specific AUC":
                        threshold_auc,

                    "AUC 95% CI lower":
                        threshold_lower,

                    "AUC 95% CI upper":
                        threshold_upper,

                    "TP":
                        metrics["TP"],

                    "TN":
                        metrics["TN"],

                    "FP":
                        metrics["FP"],

                    "FN":
                        metrics["FN"],

                    "Sensitivity":
                        metrics["Sensitivity"],

                    "Specificity":
                        metrics["Specificity"],

                    "PPV":
                        metrics["PPV"],

                    "NPV":
                        metrics["NPV"],

                    "Accuracy":
                        metrics["Accuracy"],

                    "Balanced Accuracy":
                        metrics["Balanced Accuracy"],

                    "Lesions":
                        len(temp),

                    "Patients":
                        temp[PATIENT_ID_COLUMN].nunique(),

                    "AUC CI method":
                        "Patient-level cluster bootstrap"
                }
            )

    auc_df = pd.DataFrame(
        auc_rows
    )

    threshold_df = pd.DataFrame(
        threshold_rows
    )

    # ========================================================
    # Lesion-level ROC plot
    # ========================================================

    plt.figure(
        figsize=(8, 6)
    )

    for idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        if score_col not in roc_objects:

            continue

        info = roc_objects[
            score_col
        ]

        label_text = (
            f"{DISPLAY_NAMES[idx]} "
            f"(AUC = {info['auc']:.3f} "
            f"[{info['lower']:.3f}-"
            f"{info['upper']:.3f}])"
        )

        plt.plot(
            info["fpr"],
            info["tpr"],
            color=COLORS[idx],
            linewidth=2,
            label=label_text
        )

    plt.plot(
        [0, 1],
        [0, 1],
        color="navy",
        linestyle="--",
        linewidth=1.2
    )

    plt.xlim(
        [-0.04, 1.04]
    )

    plt.ylim(
        [-0.04, 1.04]
    )

    plt.xlabel(
        "1 - Specificity"
    )

    plt.ylabel(
        "Sensitivity"
    )

    plt.title(
        f"Lesion-level ROC - {dataset_name}"
    )

    plt.legend(
        loc="lower right"
    )

    plt.grid(
        False
    )

    safe_name = (
        str(dataset_name)
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
        .replace(":", "_")
    )

    for ext in [
        "pdf",
        "svg",
        "png"
    ]:

        file_name = os.path.join(
            output_dir,
            f"{safe_name}-Lesion-level-ROC.{ext}"
        )

        if ext == "png":

            plt.savefig(
                file_name,
                dpi=DPI,
                bbox_inches="tight"
            )

        else:

            plt.savefig(
                file_name,
                format=ext,
                bbox_inches="tight"
            )

    plt.close()

    return (
        auc_df,
        threshold_df
    )


# ============================================================
# 13. Patient-level analysis
# ============================================================

def analyze_patient_level_dataset(
    df,
    dataset_name,
    output_dir
):

    print("\n" + "#" * 72)

    print(
        f"PATIENT-LEVEL：{dataset_name}"
    )

    print("#" * 72)

    patient_df = build_patient_level_data(
        df=df,
        dataset_name=dataset_name
    )

    if len(patient_df) == 0:

        return (
            pd.DataFrame(),
            pd.DataFrame()
        )

    if patient_df[
        OUTCOME_COLUMN
    ].nunique() < 2:

        return (
            pd.DataFrame(),
            pd.DataFrame()
        )

    auc_rows = []

    threshold_rows = []

    roc_objects = {}

    for idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        display_name = DISPLAY_NAMES[
            idx
        ]

        temp = patient_df[
            [
                PATIENT_ID_COLUMN,
                OUTCOME_COLUMN,
                score_col
            ]
        ].dropna().copy()

        if len(temp) == 0:

            continue

        y_true = temp[
            OUTCOME_COLUMN
        ].astype(int).to_numpy()

        y_score = temp[
            score_col
        ].astype(float).to_numpy()

        if np.unique(y_true).size < 2:

            continue

        # ----------------------------------------------------
        # Ordinal AUC
        # ----------------------------------------------------

        fpr, tpr, _ = roc_curve(
            y_true,
            y_score
        )

        (
            auc_value,
            lower,
            upper
        ) = bootstrap_auc_standard(
            y_true=y_true,
            y_score=y_score,
            n_bootstrap=N_BOOTSTRAP,
            seed=RANDOM_SEED + idx
        )

        roc_objects[
            score_col
        ] = {
            "fpr": fpr,
            "tpr": tpr,
            "auc": auc_value,
            "lower": lower,
            "upper": upper
        }

        auc_rows.append(
            {
                "Dataset":
                    dataset_name,

                "Analysis level":
                    "Patient",

                "Variable":
                    score_col,

                "Model":
                    display_name,

                "AUC":
                    auc_value,

                "95% CI lower":
                    lower,

                "95% CI upper":
                    upper,

                "Patients":
                    len(temp),

                "CI method":
                    "Patient-level bootstrap"
            }
        )

        print(
            f"\n{display_name}: "
            f"AUC={auc_value:.3f} "
            f"({lower:.3f}–{upper:.3f})"
        )

        # ----------------------------------------------------
        # >=4 and >=3
        # ----------------------------------------------------

        for threshold in [
            4,
            3
        ]:

            metrics = diagnostic_metrics(
                y_true=y_true,
                score=y_score,
                threshold=threshold
            )

            (
                threshold_auc,
                threshold_lower,
                threshold_upper
            ) = threshold_auc_patient_bootstrap(
                y_true=y_true,
                score=y_score,
                threshold=threshold,
                n_bootstrap=N_BOOTSTRAP,
                seed=(
                    RANDOM_SEED
                    + idx * 100
                    + threshold
                )
            )

            threshold_rows.append(
                {
                    "Dataset":
                        dataset_name,

                    "Analysis level":
                        "Patient",

                    "Variable":
                        score_col,

                    "Model":
                        display_name,

                    "Threshold":
                        f">={threshold}",

                    "Threshold-specific AUC":
                        threshold_auc,

                    "AUC 95% CI lower":
                        threshold_lower,

                    "AUC 95% CI upper":
                        threshold_upper,

                    "TP":
                        metrics["TP"],

                    "TN":
                        metrics["TN"],

                    "FP":
                        metrics["FP"],

                    "FN":
                        metrics["FN"],

                    "Sensitivity":
                        metrics["Sensitivity"],

                    "Specificity":
                        metrics["Specificity"],

                    "PPV":
                        metrics["PPV"],

                    "NPV":
                        metrics["NPV"],

                    "Accuracy":
                        metrics["Accuracy"],

                    "Balanced Accuracy":
                        metrics["Balanced Accuracy"],

                    "Patients":
                        len(temp),

                    "AUC CI method":
                        "Patient-level bootstrap"
                }
            )

    auc_df = pd.DataFrame(
        auc_rows
    )

    threshold_df = pd.DataFrame(
        threshold_rows
    )

    # ========================================================
    # Patient ROC
    # ========================================================

    plt.figure(
        figsize=(8, 6)
    )

    for idx, score_col in enumerate(
        SCORE_COLUMNS
    ):

        if score_col not in roc_objects:

            continue

        info = roc_objects[
            score_col
        ]

        label_text = (
            f"{DISPLAY_NAMES[idx]} "
            f"(AUC = {info['auc']:.3f} "
            f"[{info['lower']:.3f}-"
            f"{info['upper']:.3f}])"
        )

        plt.plot(
            info["fpr"],
            info["tpr"],
            color=COLORS[idx],
            linewidth=2,
            label=label_text
        )

    plt.plot(
        [0, 1],
        [0, 1],
        color="navy",
        linestyle="--",
        linewidth=1.2
    )

    plt.xlim(
        [-0.04, 1.04]
    )

    plt.ylim(
        [-0.04, 1.04]
    )

    plt.xlabel(
        "1 - Specificity"
    )

    plt.ylabel(
        "Sensitivity"
    )

    plt.title(
        f"Patient-level ROC - {dataset_name}"
    )

    plt.legend(
        loc="lower right"
    )

    plt.grid(
        False
    )

    safe_name = (
        str(dataset_name)
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
        .replace(":", "_")
    )

    for ext in [
        "pdf",
        "svg",
        "png"
    ]:

        file_name = os.path.join(
            output_dir,
            f"{safe_name}-Patient-level-ROC.{ext}"
        )

        if ext == "png":

            plt.savefig(
                file_name,
                dpi=DPI,
                bbox_inches="tight"
            )

        else:

            plt.savefig(
                file_name,
                format=ext,
                bbox_inches="tight"
            )

    plt.close()

    return (
        auc_df,
        threshold_df
    )


# ============================================================
# 14. AUC text
# ============================================================

def add_auc_ci_text(
    df,
    auc_column="AUC",
    lower_column="95% CI lower",
    upper_column="95% CI upper",
    output_column="AUC (95% CI)"
):

    if len(df) == 0:

        return df

    df = df.copy()

    df[
        output_column
    ] = df.apply(
        lambda row:
        (
            f"{row[auc_column]:.3f} "
            f"({row[lower_column]:.3f}–"
            f"{row[upper_column]:.3f})"
        )
        if (
            pd.notna(
                row[auc_column]
            )
            and pd.notna(
                row[lower_column]
            )
            and pd.notna(
                row[upper_column]
            )
        )
        else "",
        axis=1
    )

    return df


# ============================================================
# 15. Excel formatting
# ============================================================

def format_excel(
    excel_file
):

    try:

        from openpyxl import load_workbook

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment
        )

        from openpyxl.utils import (
            get_column_letter
        )

        wb = load_workbook(
            excel_file
        )

        for ws in wb.worksheets:

            for cell in ws[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )

                cell.fill = PatternFill(
                    "solid",
                    fgColor="4F81BD"
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            for column_cells in ws.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    value = (
                        ""
                        if cell.value is None
                        else str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        len(value)
                    )

                ws.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    38
                )

            ws.freeze_panes = "A2"

            ws.auto_filter.ref = (
                ws.dimensions
            )

            header_map = {
                cell.value: cell.column
                for cell in ws[1]
            }

            decimal_columns = [
                "AUC",
                "95% CI lower",
                "95% CI upper",
                "Threshold-specific AUC",
                "AUC 95% CI lower",
                "AUC 95% CI upper",
                "Sensitivity",
                "Specificity",
                "PPV",
                "NPV",
                "Accuracy",
                "Balanced Accuracy"
            ]

            for col_name in decimal_columns:

                if col_name not in header_map:

                    continue

                col_idx = header_map[
                    col_name
                ]

                for row_idx in range(
                    2,
                    ws.max_row + 1
                ):

                    ws.cell(
                        row=row_idx,
                        column=col_idx
                    ).number_format = "0.000"

        wb.save(
            excel_file
        )

    except Exception as e:

        print(
            f"Excel 格式化跳过：{e}"
        )


# ============================================================
# 16. Main
# ============================================================

def main():

    root = Tk()

    root.withdraw()

    try:

        root.attributes(
            "-topmost",
            True
        )

    except Exception:

        pass

    print("\n" + "=" * 72)

    print(
        "VI-RADS lesion-level + patient-level ROC analysis"
    )

    print("=" * 72)

    print(
        "\n包括："
    )

    print(
        "1. VI-RADS 1–5 ordinal AUC"
    )

    print(
        "2. VI-RADS >=4 threshold-specific AUC"
    )

    print(
        "3. VI-RADS >=3 threshold-specific AUC"
    )

    print(
        "4. Lesion-level patient-cluster bootstrap"
    )

    print(
        "5. Patient-level ordinary bootstrap"
    )

    print(
        "6. 不使用 Youden cutoff"
    )

    file_path = filedialog.askopenfilename(
        title="选择 VI-RADS 数据文件",
        filetypes=[
            (
                "Excel / CSV",
                "*.xlsx *.xls *.csv"
            ),
            (
                "All files",
                "*.*"
            )
        ]
    )

    root.destroy()

    if not file_path:

        print(
            "没有选择文件。"
        )

        return

    df = load_data(
        file_path
    )

    print(
        f"\n读取文件：{file_path}"
    )

    print(
        f"数据维度：{df.shape}"
    )

    # --------------------------------------------------------
    # 必要列
    # --------------------------------------------------------

    required = [
        COHORT_COLUMN,
        PATIENT_ID_COLUMN,
        OUTCOME_COLUMN
    ] + SCORE_COLUMNS

    missing = [
        col
        for col in required
        if col not in df.columns
    ]

    if missing:

        raise ValueError(
            f"缺少列：{missing}"
        )

    # --------------------------------------------------------
    # cohort cleanup
    # --------------------------------------------------------

    mask = df[
        COHORT_COLUMN
    ].notna()

    df.loc[
        mask,
        COHORT_COLUMN
    ] = (
        df.loc[
            mask,
            COHORT_COLUMN
        ]
        .astype(str)
        .str.strip()
        .str.lower()
    )

    df.loc[
        df[COHORT_COLUMN] == "",
        COHORT_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # Patient ID cleanup
    # --------------------------------------------------------

    mask = df[
        PATIENT_ID_COLUMN
    ].notna()

    df.loc[
        mask,
        PATIENT_ID_COLUMN
    ] = (
        df.loc[
            mask,
            PATIENT_ID_COLUMN
        ]
        .astype(str)
        .str.strip()
    )

    df.loc[
        df[PATIENT_ID_COLUMN] == "",
        PATIENT_ID_COLUMN
    ] = np.nan

    # --------------------------------------------------------
    # Numeric conversion
    # --------------------------------------------------------

    df[
        OUTCOME_COLUMN
    ] = pd.to_numeric(
        df[
            OUTCOME_COLUMN
        ],
        errors="coerce"
    )

    for score_col in SCORE_COLUMNS:

        df[
            score_col
        ] = pd.to_numeric(
            df[
                score_col
            ],
            errors="coerce"
        )

    # --------------------------------------------------------
    # cohorts
    # --------------------------------------------------------

    cohort_names = (
        df[
            COHORT_COLUMN
        ]
        .dropna()
        .unique()
        .tolist()
    )

    if len(cohort_names) == 0:

        raise ValueError(
            "cohort 列没有有效值。"
        )

    print(
        "\nCohorts:"
    )

    print(
        df[
            COHORT_COLUMN
        ]
        .value_counts(
            dropna=False
        )
    )

    # --------------------------------------------------------
    # output
    # --------------------------------------------------------

    input_dir = os.path.dirname(
        os.path.abspath(
            file_path
        )
    )

    output_dir = os.path.join(
        input_dir,
        "Direct_VIRADS_ROC_results"
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    all_lesion_auc = []

    all_lesion_threshold = []

    all_patient_auc = []

    all_patient_threshold = []

    # ========================================================
    # Cohort loop
    # ========================================================

    for cohort_name in cohort_names:

        df_current = df[
            df[
                COHORT_COLUMN
            ] == cohort_name
        ].copy()

        (
            lesion_auc_df,
            lesion_threshold_df
        ) = analyze_lesion_level_dataset(
            df=df_current,
            dataset_name=str(
                cohort_name
            ),
            output_dir=output_dir
        )

        if len(
            lesion_auc_df
        ) > 0:

            all_lesion_auc.append(
                lesion_auc_df
            )

        if len(
            lesion_threshold_df
        ) > 0:

            all_lesion_threshold.append(
                lesion_threshold_df
            )

        (
            patient_auc_df,
            patient_threshold_df
        ) = analyze_patient_level_dataset(
            df=df_current,
            dataset_name=str(
                cohort_name
            ),
            output_dir=output_dir
        )

        if len(
            patient_auc_df
        ) > 0:

            all_patient_auc.append(
                patient_auc_df
            )

        if len(
            patient_threshold_df
        ) > 0:

            all_patient_threshold.append(
                patient_threshold_df
            )

    # ========================================================
    # Combine
    # ========================================================

    lesion_auc = pd.concat(
        all_lesion_auc,
        ignore_index=True
    )

    lesion_threshold = pd.concat(
        all_lesion_threshold,
        ignore_index=True
    )

    patient_auc = pd.concat(
        all_patient_auc,
        ignore_index=True
    )

    patient_threshold = pd.concat(
        all_patient_threshold,
        ignore_index=True
    )

    # --------------------------------------------------------
    # AUC text
    # --------------------------------------------------------

    lesion_auc = add_auc_ci_text(
        lesion_auc
    )

    patient_auc = add_auc_ci_text(
        patient_auc
    )

    lesion_threshold = add_auc_ci_text(
        lesion_threshold,
        auc_column="Threshold-specific AUC",
        lower_column="AUC 95% CI lower",
        upper_column="AUC 95% CI upper",
        output_column="Threshold AUC (95% CI)"
    )

    patient_threshold = add_auc_ci_text(
        patient_threshold,
        auc_column="Threshold-specific AUC",
        lower_column="AUC 95% CI lower",
        upper_column="AUC 95% CI upper",
        output_column="Threshold AUC (95% CI)"
    )

    # --------------------------------------------------------
    # >=4 / >=3
    # --------------------------------------------------------

    lesion_ge4 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    lesion_ge3 = (
        lesion_threshold[
            lesion_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge4 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=4"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    patient_ge3 = (
        patient_threshold[
            patient_threshold[
                "Threshold"
            ] == ">=3"
        ]
        .copy()
        .reset_index(
            drop=True
        )
    )

    # ========================================================
    # Excel
    # ========================================================

    excel_file = os.path.join(
        output_dir,
        "VI-RADS_ROC_lesion_and_patient_level.xlsx"
    )

    with pd.ExcelWriter(
        excel_file,
        engine="openpyxl"
    ) as writer:

        lesion_auc.to_excel(
            writer,
            sheet_name="Lesion_AUC",
            index=False
        )

        lesion_ge4.to_excel(
            writer,
            sheet_name="Lesion_ge4",
            index=False
        )

        lesion_ge3.to_excel(
            writer,
            sheet_name="Lesion_ge3",
            index=False
        )

        patient_auc.to_excel(
            writer,
            sheet_name="Patient_AUC",
            index=False
        )

        patient_ge4.to_excel(
            writer,
            sheet_name="Patient_ge4",
            index=False
        )

        patient_ge3.to_excel(
            writer,
            sheet_name="Patient_ge3",
            index=False
        )

    format_excel(
        excel_file
    )

    # ========================================================
    # CSV
    # ========================================================

    output_tables = {
        "Lesion_level_AUC.csv":
            lesion_auc,

        "Lesion_level_VIRADS_ge4.csv":
            lesion_ge4,

        "Lesion_level_VIRADS_ge3.csv":
            lesion_ge3,

        "Patient_level_AUC.csv":
            patient_auc,

        "Patient_level_VIRADS_ge4.csv":
            patient_ge4,

        "Patient_level_VIRADS_ge3.csv":
            patient_ge3
    }

    for filename, table in output_tables.items():

        table.to_csv(
            os.path.join(
                output_dir,
                filename
            ),
            index=False,
            encoding="utf-8-sig"
        )

    # ========================================================
    # Console summary
    # ========================================================

    print("\n" + "=" * 72)

    print(
        "LESION-LEVEL >=4"
    )

    print("=" * 72)

    print(
        lesion_ge4[
            [
                "Dataset",
                "Model",
                "Threshold-specific AUC",
                "AUC 95% CI lower",
                "AUC 95% CI upper",
                "Sensitivity",
                "Specificity",
                "PPV",
                "NPV",
                "Accuracy",
                "Balanced Accuracy"
            ]
        ].to_string(
            index=False
        )
    )

    print("\n" + "=" * 72)

    print(
        "LESION-LEVEL >=3"
    )

    print("=" * 72)

    print(
        lesion_ge3[
            [
                "Dataset",
                "Model",
                "Threshold-specific AUC",
                "AUC 95% CI lower",
                "AUC 95% CI upper",
                "Sensitivity",
                "Specificity",
                "PPV",
                "NPV",
                "Accuracy",
                "Balanced Accuracy"
            ]
        ].to_string(
            index=False
        )
    )

    print("\n" + "=" * 72)

    print(
        "PATIENT-LEVEL >=4"
    )

    print("=" * 72)

    print(
        patient_ge4[
            [
                "Dataset",
                "Model",
                "Threshold-specific AUC",
                "AUC 95% CI lower",
                "AUC 95% CI upper",
                "Sensitivity",
                "Specificity",
                "PPV",
                "NPV",
                "Accuracy",
                "Balanced Accuracy"
            ]
        ].to_string(
            index=False
        )
    )

    print("\n" + "=" * 72)

    print(
        "PATIENT-LEVEL >=3"
    )

    print("=" * 72)

    print(
        patient_ge3[
            [
                "Dataset",
                "Model",
                "Threshold-specific AUC",
                "AUC 95% CI lower",
                "AUC 95% CI upper",
                "Sensitivity",
                "Specificity",
                "PPV",
                "NPV",
                "Accuracy",
                "Balanced Accuracy"
            ]
        ].to_string(
            index=False
        )
    )

    print("\n" + "=" * 72)

    print(
        "分析完成"
    )

    print("=" * 72)

    print(
        f"\n结果目录：\n{output_dir}"
    )

    print(
        f"\nExcel：\n{excel_file}"
    )

    print(
        "\nExcel sheets:"
    )

    print(
        "Lesion_AUC"
    )

    print(
        "Lesion_ge4"
    )

    print(
        "Lesion_ge3"
    )

    print(
        "Patient_AUC"
    )

    print(
        "Patient_ge4"
    )

    print(
        "Patient_ge3"
    )


# ============================================================
# 17. Run
# ============================================================

if __name__ == "__main__":

    try:

        main()

    except Exception as e:

        print("\n" + "=" * 72)

        print(
            "程序发生错误"
        )

        print("=" * 72)

        print(
            f"\n{type(e).__name__}: {e}"
        )

        print(
            "\n请检查："
        )

        print(
            "1. Group 是否为 0/1"
        )

        print(
            "2. PatientID 是否存在"
        )

        print(
            "3. cohort 是否存在"
        )

        print(
            "4. SCORE_COLUMNS 列名是否完全一致"
        )

        print(
            "5. VI-RADS 是否为 1–5"
        )

        print(
            "6. 每个 cohort 是否同时有 MIBC 和 NMIBC"
        )

    finally:

        try:

            input(
                "\n按 Enter 键退出..."
            )

        except Exception:

            pass